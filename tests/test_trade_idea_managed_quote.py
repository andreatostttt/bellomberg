"""Quoted managed-care bindings share the precision of the live ROUND formula."""
from openpyxl import load_workbook

from bellomberg.valuation.dcf_engine import generate_valuation
from bellomberg.valuation.trade_idea_model import model_exhibits
from test_managed_care_integration import make_bundle


def test_managed_care_quoted_rounding_matches_common_formula_and_cache(tmp_path):
    bundle = make_bundle(kind='FY', count=10)
    payload = generate_valuation(bundle['case']['ticker'], prepared_bundle=bundle, output_dir=str(tmp_path))
    assert payload['valuation_usability']['usable']
    quote = next(row['value'] for row in bundle['case']['records'] if row['driver'] == 'quotation')
    factor = quote['financial_to_quote_rate'] * quote['quote_units_per_currency'] * quote['shares_per_quote']
    formulas = load_workbook(payload['path'], data_only=False, read_only=True)
    cache = load_workbook(payload['path'], data_only=True, read_only=True)
    try:
        for index, (scenario, coordinate) in enumerate(zip(('bear', 'base', 'bull'), ('D8', 'E8', 'F8'))):
            raw = payload['managed_care']['scenarios'][scenario]['fair_value_per_share']
            assert payload['fair_value_' + scenario] == round(raw * factor, 2)
            assert 'ROUND(' in formulas['Summary'][coordinate].value and ',2)' in formulas['Summary'][coordinate].value
            assert cache['Model Checks']['C' + str(index + 8)].value == raw
    finally:
        formulas.close(); cache.close()
    assert payload['managed_care']['scenarios']['base']['fair_value_per_share'] != payload['fair_value_base']
    packet = model_exhibits(payload)
    assert packet['status'] == 'complete', packet['reasons']
    values = next(row for row in packet['exhibits'] if row['id'] == 'scenarios')
    assert values['rows'][1][1] == payload['fair_value_base']
