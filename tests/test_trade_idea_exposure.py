"""Observed exposure is a separate analysis objective, never a surrogate FV."""
from copy import deepcopy
from hashlib import sha256
import json

import pytest

from bellomberg.valuation.trade_idea_model import qualification, prepare, model_exhibits, candidate_model_usability
from bellomberg.valuation.trade_idea_exposure import SCHEMA
from test_trade_idea_economic import synthetic_observed_quote_info

DAY = '2026-09-10'


def sourced_exposure(tmp_path):
    ticker = 'SYNTH-ETF'
    profile = {'status': 'ok', 'as_of': DAY, 'retrieved_at': DAY+'T12:00:00+00:00', 'source_id': 'synthetic ETF prospectus',
        'data': {'info': {'symbol': ticker, 'quoteType': 'ETF', 'shortName': 'Synthetic fund', 'currency': 'EUR', 'exchange': 'TEST'},
                 'vehicle_registry': None,
                 'evidence': [{'field': 'instrument', 'value': 'etf', 'source_id': 'synthetic', 'as_of': DAY}]}}
    identity = {'ticker': ticker, 'name': 'Synthetic fund', 'exchange': 'TEST', 'currency': 'EUR', 'status': 'confirmed'}
    profile['data']['info'].update(synthetic_observed_quote_info(ticker, DAY, price=100.))
    values = {
        'perimeter': {'entity': 'Synthetic fund', 'currency': 'EUR', 'share_class': 'Accumulating'},
        'calendar': {'valuation_date': DAY, 'periods': [], 'discount_convention': 'snapshot'},
        'identity': {'ticker': ticker, 'name': 'Synthetic fund', 'instrument': 'etf', 'issuer': 'Synthetic issuer',
                     'legal_structure': 'UCITS fund', 'underlying': 'Synthetic diversified index'},
        'holdings': {'basis': 'net_asset_exposure', 'positions': [
            {'id': 'A', 'name': 'Synthetic A', 'weight': .6, 'asset_class': 'equity', 'country': 'IT', 'currency': 'EUR'},
            {'id': 'B', 'name': 'Synthetic B', 'weight': .4, 'asset_class': 'equity', 'country': 'FR', 'currency': 'EUR'}]},
        'terms': {'replication': 'physical_full', 'benchmark': 'Synthetic diversified index',
                  'gross_exposure': 1., 'net_exposure': 1., 'leverage': 1., 'effective_date': DAY},
        'fees': {'annual_expense_ratio': .002, 'other_annual_fees_ratio': 0.},
        'risks': {'issuer': 'Observed fund issuer terms', 'counterparty': 'No swap counterparties in stated physical scope',
                  'custody': 'Observed custody terms', 'liquidity': 'Published redemption terms', 'tracking': 'Published tracking risks'},
        'liquidity': {'daily_volume': 10000., 'volume_unit': 'quoted units', 'spread_ratio': .001,
                      'observed_date': DAY, 'provider': 'Synthetic exchange observation'},
        'quotation': {'financial_currency': 'EUR', 'quote_currency': 'EUR', 'quote_unit': 'EUR',
                      'quote_units_per_currency': 1., 'financial_to_quote_rate': 1., 'shares_per_quote': 1.,
                      'share_class': 'Accumulating', 'price': 100., 'price_as_of': DAY},
    }
    observations = []
    plan = {'model': {}, 'scenarios': {}, 'analysis_rationale': 'Synthetic sourced observational exposure proof'}
    for driver, value in values.items():
        field, unit, basis, *_ = SCHEMA[driver]
        entry = {'value': deepcopy(value), 'kind': 'analyst_estimate' if driver in ('perimeter', 'calendar') else 'historical',
                 'evidence_ids': ['synthetic-public'], 'rationale': 'Explicit synthetic source observation',
                 'valid_until': DAY, 'valid_until_basis': {'policy': 'same_day', 'as_of': DAY}}
        if entry['kind'] == 'historical':
            observations.append({'field': field, 'driver': driver, 'value': deepcopy(value), 'entity': 'Synthetic fund',
                                 'period': DAY, 'unit': unit, 'accounting_basis': basis})
            entry['record_pointer'] = '/observations/' + str(len(observations)-1)
        plan['model'][driver] = entry
    text = json.dumps({'synthetic': True, 'observations': observations}, sort_keys=True, separators=(',', ':'))
    documents = [{'id': 'synthetic-public', 'url': 'https://example.org/synthetic/ETF-prospectus-and-holdings',
                  'published_at': DAY, 'text': text, 'sha256': sha256(text.encode()).hexdigest()}]
    q = qualification(ticker, identity, DAY, archive_root=tmp_path,
        providers={'profile': lambda *_a, **_k: deepcopy(profile)}, source_report={'documents': documents, 'source_plan': plan})
    return q, plan


def test_observational_exposure_has_exact_workbook_and_no_fv(tmp_path):
    q, plan = sourced_exposure(tmp_path)
    assert q['status'] == 'qualified', q['reasons']
    payload = prepare(q, lambda *_: deepcopy(plan), tmp_path / 'model')
    assert payload['analysis_usability']['usable'], payload.get('error')
    assert not payload['valuation_usability']['usable']
    assert all(payload.get('fair_value_' + scenario) is None for scenario in ('bear', 'base', 'bull'))
    assert candidate_model_usability(payload) == {'usable': True, 'kind': 'exposure', 'reasons': [],
        'objective': 'observed_exposure_analysis', 'intrinsic_value_applicable': False}
    packet = model_exhibits(payload)
    assert packet['status'] == 'complete', packet['reasons']
    assert {row['id'] for row in packet['exhibits']} == {'exposure_metrics', 'holdings'}
    assert packet['headline_values']['price']['value'] == 100.
    from openpyxl import load_workbook
    workbook = load_workbook(payload['path'], data_only=False)
    assert workbook['Summary']['D8'].data_type == 'f'
    workbook.close()
    from pathlib import Path
    sidecar = json.loads(Path(payload['path']).with_suffix('.payload.json').read_text(encoding='utf-8'))
    assert candidate_model_usability(sidecar)['usable']
    assert model_exhibits(sidecar, workbook_path=payload['path']) == packet
    durable = json.loads(json.dumps(payload, sort_keys=True))
    assert model_exhibits(durable, workbook_path=payload['path']) == packet
    wrong = deepcopy(payload); wrong['fair_value_base'] = 100.
    assert not candidate_model_usability(wrong)['usable']


def test_observed_missing_position_is_not_zero_or_normalized(tmp_path):
    q, plan = sourced_exposure(tmp_path)
    plan['model']['holdings']['value']['positions'][0]['weight'] = .5
    payload = prepare(q, lambda *_: deepcopy(plan), tmp_path / 'missing')
    assert not payload.get('analysis_usability', {}).get('usable')
    assert not candidate_model_usability(payload)['usable']
    assert any('holdings' in row['reason'] for row in payload['preparation']['issues'])


def test_exposure_public_projection_reproves_observed_text_and_regenerates(tmp_path):
    from bellomberg.valuation.trade_idea_model import prepare_shareable_model, public_model_payload
    q, plan = sourced_exposure(tmp_path)
    private = prepare(q, lambda *_: deepcopy(plan), tmp_path / 'private')
    canary = 'PERSONAL_EXPOSURE_PM_CANARY_79341'
    private['preparation']['provenance']['personal_commentary'] = canary
    public = public_model_payload(private)
    assert public['status'] == 'ready', public['reasons']
    assert canary not in json.dumps(public)
    shared = prepare_shareable_model(private, tmp_path / 'shared')
    assert candidate_model_usability(shared)['usable']
    assert not shared['valuation_usability']['usable']
    assert model_exhibits(shared)['status'] == 'complete'
    assert shared['snapshot_id'] != private['snapshot_id']
    assert shared['generation_id'] != private['generation_id']


@pytest.mark.parametrize('sheet,cell,value', [
    ('Exposure', 'D8', .5),
    ('Exposure', 'C8', 'Unbound holding name'),
    ('Exposure', 'H8', '=ABS(D9)'),
    ('Summary', 'D8', "=SUM('Exposure'!H8:H8)"),
    ('Summary', 'D13', .002),
    ('Summary', 'D15', 101.),
    ('Terms and risks', 'C20', .003),
])
def test_exposure_changed_workbook_is_rejected_even_with_updated_hash(tmp_path, sheet, cell, value):
    from pathlib import Path
    from openpyxl import load_workbook
    q, plan = sourced_exposure(tmp_path)
    payload = prepare(q, lambda *_: deepcopy(plan), tmp_path / 'model')
    path = Path(payload['path'])
    workbook = load_workbook(path)
    workbook[sheet][cell] = value
    workbook.save(path); workbook.close()
    payload['workbook_sha256'] = sha256(path.read_bytes()).hexdigest()
    packet = model_exhibits(payload)
    assert packet['status'] == 'blocked'
    assert packet['reasons']


def test_exposure_tampered_input_binding_is_rejected(tmp_path):
    q, plan = sourced_exposure(tmp_path)
    payload = prepare(q, lambda *_: deepcopy(plan), tmp_path / 'model')
    payload['calculation_details']['workbook_bindings']['input_refs'][0]['cell'] = "'Exposure'!D9"
    assert model_exhibits(payload)['status'] == 'blocked'
