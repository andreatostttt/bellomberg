"""Real workbook exports bind readable sources to records consumed by the compiler."""
from copy import deepcopy
from hashlib import sha256
import json

from openpyxl import Workbook, load_workbook

from test_input_preparation import _bundle, _documents, _operating_plan
from test_sector_operating_drivers import bundle_for
from bellomberg.valuation.dcf_engine import generate_valuation
from bellomberg.valuation.preparation_service import prepare_and_generate
from bellomberg.valuation.sourcebook_presentation import present_sourcebook


def _prepared(tmp_path):
    plan = _operating_plan()
    text = json.dumps({'facts': [{'taxonomy': 'us-gaap', 'concept': 'Revenues',
        'value': 100000000, 'unit': 'EUR', 'start': '2025-01-01', 'end': '2025-12-31'}]})
    documents = _documents() + [{'id': 'primary-facts', 'text': text,
        'sha256': sha256(text.encode()).hexdigest(),
        'url': 'https://example.org/issuer/facts.json', 'published_at': '2026-02-01'}]
    item = plan['model']['historical_revenue']
    for key in ('evidence_quote', 'quoted_value', 'quoted_unit', 'period_quote'):
        item.pop(key)
    item.update(evidence_ids=['primary-facts'], quoted_value=100000000, quoted_unit='EUR',
        evidence_pointer={'value': '/facts/0/value', 'unit': '/facts/0/unit', 'period': '/facts/0/end'})
    before = deepcopy((plan, documents))
    result = prepare_and_generate(_bundle(), documents=documents, propose=lambda *_: plan,
                                  output_dir=tmp_path)
    assert (plan, documents) == before
    assert result['valuation_usability']['usable'], result.get('error')
    return result


def _rows(sheet, scenario, driver):
    return [row for row in sheet.iter_rows(min_row=8) if row[1].value == scenario and row[2].value == driver]


def test_real_compiler_exports_consumed_source_ids_dates_pointers_and_rationale(tmp_path):
    result = _prepared(tmp_path)
    wb = load_workbook(result['path'], data_only=False)
    assert 'Sources' in wb.sheetnames
    rows = _rows(wb['Sources'], 'model', 'historical_revenue')
    assert len(rows) == 1
    row = rows[0]
    assert row[3].value == 'Fatto riportato'
    assert row[4].value == '100.0'
    assert 'EUR million' in row[5].value and '2025-12-31' in row[5].value
    assert row[7].value == 'primary-facts'
    assert '2026-02-01' in row[8].value
    assert row[9].value == 'https://example.org/issuer/facts.json'
    assert row[9].hyperlink.target == row[9].value
    assert '/facts/0/value' in row[10].value and '100000000' in row[10].value
    assert row[11].value == 'Record consumato; prova del piano collegata'
    estimate = _rows(wb['Sources'], 'base', 'revenue_growth')[0]
    assert estimate[3].value == 'Stima analista'
    assert estimate[6].value == result['preparation']['proposal']['plan']['scenarios']['base']['revenue_growth']['rationale']
    assert wb['Assumptions']['B8'].hyperlink.location == f"'Sources'!B{estimate[1].row}"
    assert wb['Sources'].protection.sheet is True
    assert not any(cell.data_type == 'f' for row in wb['Sources'] for cell in row)
    # Read-only evidence does not replace or strip the live calculation formulas.
    assert wb['base']['D8'].data_type == 'f'
    assert result['calculation_details']['workbook_bindings']['contract'] == 'common_workbook_bindings/1'
    wb.close()


def test_approved_record_guidance_is_labeled_without_inventing_document_metadata(tmp_path):
    bundle = bundle_for()
    records = deepcopy(bundle['case']['records'])
    item = next(row for row in records if row['scenario'] == 'base' and row['driver'] == 'revenue_growth')
    item['kind'] = 'company_guidance'
    item['rationale'] = 'Explicitly declared guidance in this frozen test record'
    # Build the frozen provider envelope once; do not duplicate it as an explicit proposal.
    bundle = bundle_for(records=records)
    result = generate_valuation('SYNTH-EXT', prepared_bundle=bundle, output_dir=tmp_path)
    assert result.get('path'), result
    wb = load_workbook(result['path'])
    row = _rows(wb['Sources'], 'base', 'revenue_growth')[0]
    assert row[3].value == 'Guidance societaria'
    assert row[6].value == item['rationale']
    assert row[7].value == 'n.d.'
    assert 'n.d.' in row[10].value and 'Metadati documento n.d.' in row[11].value
    wb.close()


def test_source_view_uses_consumption_receipts_and_preserves_missing_provenance(tmp_path):
    payload = _prepared(tmp_path)
    payload['acquisition_snapshot']['case']['records'].append({'scenario': 'bull', 'driver': 'UNUSED_SENTINEL'})
    payload['preparation']['provenance']['documents'].pop('primary-facts')
    before = deepcopy(payload)
    wb = Workbook()
    assert present_sourcebook(wb, payload)
    assert payload == before
    assert not any(cell.value == 'UNUSED_SENTINEL' for row in wb['Sources'] for cell in row)
    row = _rows(wb['Sources'], 'model', 'historical_revenue')[0]
    assert row[7].value == 'primary-facts'
    assert 'Metadati documento n.d.' in row[11].value
    assert 'Pubblicazione: n.d.' in row[8].value


def test_old_plan_pointer_is_not_attributed_to_a_changed_consumed_record(tmp_path):
    payload = _prepared(tmp_path)
    old = next(row for row in payload['preparation']['proposal']['method_records']
               if row['scenario'] == 'model' and row['driver'] == 'historical_revenue')
    old['value'] = 999.0
    wb = Workbook()
    present_sourcebook(wb, payload)
    row = _rows(wb['Sources'], 'model', 'historical_revenue')[0]
    assert '/facts/0/value' not in row[10].value
    assert 'Piano non collegato al record consumato' in row[11].value


def test_changed_plan_pointer_cannot_claim_the_original_compilation_receipt(tmp_path):
    payload = _prepared(tmp_path)
    payload['preparation']['proposal']['plan']['model']['historical_revenue']['evidence_pointer']['value'] = '/facts/999/value'
    wb = Workbook()
    present_sourcebook(wb, payload)
    row = _rows(wb['Sources'], 'model', 'historical_revenue')[0]
    assert '/facts/999/value' not in row[10].value
    assert 'Piano non collegato al record consumato' in row[11].value


def test_changed_provenance_is_exposed_without_changing_or_relabeling_model_inputs(tmp_path):
    payload = _prepared(tmp_path)
    records = deepcopy(payload['acquisition_snapshot']['case']['records'])
    payload['preparation']['provenance']['documents']['primary-facts']['url'] = 'https://wrong.example/document'
    wb = Workbook()
    present_sourcebook(wb, payload)
    row = _rows(wb['Sources'], 'model', 'historical_revenue')[0]
    assert row[9].value == 'https://example.org/issuer/facts.json'
    assert 'Metadati documento incoerenti' in row[11].value
    assert payload['acquisition_snapshot']['case']['records'] == records


def test_receipt_mismatch_is_visible_and_does_not_borrow_another_record(tmp_path):
    payload = _prepared(tmp_path)
    receipt = next(row for row in payload['input_consumption']['consumed_records']
                   if row['scenario'] == 'model' and row['driver'] == 'historical_revenue')
    receipt['record_index'] = 0
    wb = Workbook()
    present_sourcebook(wb, payload)
    row = _rows(wb['Sources'], 'model', 'historical_revenue')[0]
    assert row[4].value == 'n.d.'
    assert 'Ricevuta di consumo non corrispondente' in row[11].value


def test_untrusted_and_long_source_text_is_literal_and_not_truncated(tmp_path):
    payload = _prepared(tmp_path)
    record = next(row for row in payload['acquisition_snapshot']['case']['records']
                  if row['scenario'] == 'base' and row['driver'] == 'revenue_growth')
    text = '=HYPERLINK("https://invalid.example")' + '\U0001f4c4' * 20000 + 'END_MARKER'
    record['rationale'] = text
    record['source_id'] = '=HYPERLINK("https://bad.example")'
    receipt = next(row for row in payload['input_consumption']['consumed_records']
                   if row['scenario'] == 'base' and row['driver'] == 'revenue_growth')
    receipt['source_id'] = record['source_id']
    wb = Workbook()
    present_sourcebook(wb, payload)
    path = tmp_path / 'literal-sources.xlsx'
    wb.save(path)
    reopened = load_workbook(path)
    sheet = reopened['Sources']
    rows = _rows(sheet, 'base', 'revenue_growth')
    assert ''.join(row[6].value or '' for row in rows) == text
    assert all(cell.data_type != 'f' for row in sheet for cell in row)
    assert any(cell.value == record['source_id'] and cell.hyperlink is None for row in sheet for cell in row)
    reopened.close()
