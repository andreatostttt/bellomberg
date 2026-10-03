"""Explicitly synthetic commodity-trust observations, never an intrinsic value."""
from copy import deepcopy
from hashlib import sha256
import json

import pytest


def values():
    from bellomberg.valuation.trade_idea_commodity_exposure import POLICY
    return {
        'policy': deepcopy(POLICY),
        'perimeter': {'entity': 'Synthetic Gold Trust', 'currency': 'USD', 'share_class': 'Shares'},
        'calendar': {'valuation_date': '2026-06-30', 'periods': [], 'discount_convention': 'snapshot'},
        'identity': {'ticker': 'SYNTH-GOLD', 'name': 'Synthetic Gold Trust', 'instrument': 'etf',
            'issuer': 'Synthetic Gold Trust', 'legal_structure': 'Commodity trust', 'underlying': 'Physical gold'},
        'holdings': {'assets': {'gold': 1002.}, 'liabilities': {'sponsor_fee_payable': 2.},
            'reported': {'total_assets': 1002., 'total_liabilities': 2., 'net_assets': 1000., 'shares': 100.},
            'as_of': '2026-06-30', 'currency': 'USD'},
        'terms': {'replication': 'direct_asset', 'benchmark': 'Published gold benchmark',
            'effective_date': '2026-06-30', 'investment_leverage': None, 'leverage_basis': 'not_separately_established'},
        'fees': {'period_start': '2026-01-01', 'period_end': '2026-06-30', 'sponsor_fees': 1.,
            'total_expenses': 1., 'annualized_expense_ratio': .0025,
            'annualization_basis': 'reported_annualized_historical_expense_ratio',
            'future_expense_ratio': None, 'other_future_fees_ratio': None,
            'contingent_expenses': 'Published contingent expenses remain uncertain'},
        'risks': {'issuer': 'Published trust risks', 'counterparty': 'Published counterparty risks',
            'custody': 'Published custody risks', 'liquidity': 'Published liquidity risks', 'tracking': 'Published tracking risks'},
        'liquidity': {'daily_volume': 1000., 'volume_unit': 'quoted units', 'spread_ratio': .0001,
            'spread_basis': '30_day_median_bid_ask', 'spread_window_days': 30,
            'observed_date': '2026-09-25', 'provider': 'Synthetic issuer observation'},
        'quotation': {'financial_currency': 'USD', 'quote_currency': 'USD', 'quote_unit': 'USD',
            'quote_units_per_currency': 1., 'financial_to_quote_rate': 1., 'shares_per_quote': 1.,
            'share_class': 'Shares', 'price': 11., 'price_as_of': '2026-09-25'},
    }


def sourced(tmp_path):
    from bellomberg.valuation.trade_idea_commodity_exposure import SCHEMA, record_period
    from bellomberg.valuation.trade_idea_model import qualification
    from test_trade_idea_economic import synthetic_observed_quote_info
    model = values(); day = '2026-09-28'; entity = model['perimeter']['entity']
    plan = {'model': {}, 'scenarios': {}, 'analysis_rationale': 'Synthetic dated historical accounting and separately dated market observations'}
    observations = []
    for driver, value in model.items():
        historical = driver not in ('perimeter', 'calendar', 'policy')
        entry = {'value': deepcopy(value), 'kind': 'historical' if historical else 'analyst_estimate',
            'evidence_ids': ['synthetic-gold'], 'rationale': 'Explicit synthetic observed contract',
            'valid_until': day, 'valid_until_basis': {'policy': 'same_day', 'as_of': day}}
        if historical:
            field, unit, basis, *_ = SCHEMA[driver]
            observations.append({'driver': driver, 'field': field, 'value': deepcopy(value), 'entity': entity,
                'period': record_period(driver, value, model['calendar']), 'unit': unit, 'accounting_basis': basis})
            entry['record_pointer'] = '/observations/' + str(len(observations)-1)
        plan['model'][driver] = entry
    text = json.dumps({'synthetic': True, 'observations': observations}, sort_keys=True)
    doc = {'id': 'synthetic-gold', 'url': 'https://example.org/synthetic/gold-trust', 'published_at': day,
        'text': text, 'sha256': sha256(text.encode()).hexdigest()}
    info = {'symbol': 'SYNTH-GOLD', 'longName': entity, 'financialCurrency': 'USD', 'quoteType': 'ETF'}
    info.update(synthetic_observed_quote_info('SYNTH-GOLD', day, currency='USD', exchange='TEST'))
    profile = {'status': 'ok', 'as_of': day, 'retrieved_at': day+'T12:00:00+00:00', 'source_id': doc['url'],
        'data': {'info': info, 'evidence': [{'field': 'instrument', 'value': 'etf', 'source_id': doc['url'], 'as_of': day}]}}
    identity = {'ticker': 'SYNTH-GOLD', 'name': entity, 'exchange': 'TEST', 'currency': 'USD', 'status': 'confirmed'}
    q = qualification('SYNTH-GOLD', identity, day, archive_root=tmp_path,
        providers={'profile': lambda *_a, **_k: deepcopy(profile)}, source_report={'documents': [doc], 'source_plan': plan})
    return q, plan


def primary_sourced(tmp_path):
    """Fictional original PDF/HTML formats; no previously prepared model/plan."""
    from test_commodity_trust_evidence import sources, normalize, ENTITY, TICKER, DAY
    from test_trade_idea_economic import synthetic_observed_quote_info
    from bellomberg.valuation.trade_idea_model import qualification
    documents = sources(tmp_path)
    normalized = normalize(documents)
    assert normalized['status'] == 'ready', normalized['issues']
    report = {'documents': documents+normalized['documents'], 'preparation_ready': False}
    info = {'symbol': TICKER, 'longName': ENTITY, 'financialCurrency': 'USD', 'quoteType': 'ETF'}
    info.update(synthetic_observed_quote_info(TICKER, DAY, currency='USD', exchange='TEST'))
    info['regularMarketPrice'] = 10.25
    profile = {'status': 'ok', 'as_of': DAY, 'retrieved_at': DAY+'T12:00:00+00:00',
        'source_id': 'https://example.org/synthetic-chart', 'data': {'info': info, 'vehicle_registry': None}}
    identity = {'ticker': TICKER, 'name': ENTITY, 'exchange': 'TEST', 'currency': 'USD', 'status': 'confirmed'}
    def no_seed(*_args, **_kwargs):
        raise AssertionError('Fresh original sources must not read a previous model')
    q = qualification(TICKER, identity, DAY, archive_root=tmp_path/'archive',
        providers={'profile': lambda *_a, **_k: deepcopy(profile)}, collector=lambda *_a, **_k: deepcopy(report),
        historical_seed_loader=no_seed)
    assert q['status'] == 'qualified', q['reasons']
    assert 'source_plan' not in report
    return q, deepcopy(q['source_report']['source_plan'])


def test_accounting_exposure_does_not_invent_investment_leverage_or_future_fee_zero():
    from bellomberg.valuation.trade_idea_commodity_exposure import project
    result = project(values(), ticker='SYNTH-GOLD', instrument='etf', cutoff='2026-09-28')
    assert result['status'] == 'complete', result['issues']
    assert result['gold_nav_ratio'] == 1.002 and result['liability_nav_ratio'] == -.002
    assert result['net_exposure'] == 1. and result['investment_leverage'] is None
    assert result['annual_declared_cost_ratio'] is None
    assert result['historical_annualized_expense_ratio'] == .0025
    assert result['liquidity']['observed_date'] != result['accounting_date']


@pytest.mark.parametrize('fault', ['ledger', 'shares', 'future_cost_zero', 'invented_leverage', 'spot_spread', 'future_liquidity', 'wrong_currency', 'wrong_quote_class'])
def test_commodity_observations_reject_false_scope_and_missing_semantics(fault):
    from bellomberg.valuation.trade_idea_commodity_exposure import project
    model = values()
    if fault == 'ledger': model['holdings']['reported']['net_assets'] += 1
    elif fault == 'shares': model['holdings']['reported']['shares'] = 0
    elif fault == 'future_cost_zero': model['fees']['other_future_fees_ratio'] = 0.
    elif fault == 'invented_leverage': model['terms']['investment_leverage'] = 1.004
    elif fault == 'spot_spread': model['liquidity']['spread_basis'] = 'spot_bid_ask'
    elif fault == 'future_liquidity': model['liquidity']['observed_date'] = '2026-09-29'
    elif fault == 'wrong_currency': model['holdings']['currency'] = 'EUR'
    else: model['quotation']['share_class'] = 'Preferred'
    result = project(model, ticker='SYNTH-GOLD', instrument='etf', cutoff='2026-09-28')
    assert result['status'] == 'incomplete' and result['issues']


def test_commodity_common_generation_has_live_accounting_formulas_and_exact_dated_exhibits(tmp_path):
    from bellomberg.valuation.trade_idea_model import prepare, candidate_model_usability, model_exhibits
    q, plan = sourced(tmp_path)
    assert q['status'] == 'qualified', q['reasons']
    payload = prepare(q, lambda *_: deepcopy(plan), tmp_path/'model')
    assert candidate_model_usability(payload)['usable'], payload.get('error')
    assert not payload['valuation_usability']['usable']
    assert all(payload.get('fair_value_'+scenario) is None for scenario in ('bear', 'base', 'bull'))
    packet = model_exhibits(payload)
    assert packet['status'] == 'complete', packet['reasons']
    assert packet['headline_values']['price']['period'] == '2026-09-25'
    assert any(row['id'] == 'commodity_accounting' for row in packet['exhibits'])
    assert model_exhibits(json.loads(json.dumps(payload, sort_keys=True))) == packet
    from openpyxl import load_workbook
    from pathlib import Path
    workbook = load_workbook(payload['path'])
    workbook['Summary']['D8'] = 1000.
    workbook.save(payload['path']); workbook.close()
    payload['workbook_sha256'] = sha256(Path(payload['path']).read_bytes()).hexdigest()
    assert model_exhibits(payload)['status'] == 'blocked'


def test_commodity_proposer_receives_observed_contract_without_legacy_future_cost_shape(tmp_path):
    from bellomberg.valuation.trade_idea_model import prepare
    from bellomberg.valuation.trade_idea_commodity_exposure import SCHEMA, CONTRACT
    q, plan = sourced(tmp_path)
    seen = []
    def proposer(dossier, contract):
        seen.append(contract)
        return deepcopy(plan)
    prepare(q, proposer, tmp_path/'prepared')
    assert seen[0]['schema'] == SCHEMA
    assert seen[0]['method_variant'] == CONTRACT
    assert 'future_expense_ratio' in seen[0]['method_contract']
    assert 'annual_expense_ratio,other_annual_fees_ratio' not in seen[0]['method_contract']


@pytest.mark.parametrize('receipt', ['valid', 'missing', 'future'])
def test_current_quote_is_separate_from_issuer_price_and_accounting_snapshot(tmp_path, receipt):
    from bellomberg.valuation.trade_idea_model import prepare, candidate_model_usability
    from bellomberg.valuation.dcf_engine import generate_valuation
    q, plan = sourced(tmp_path)
    original = prepare(q, lambda *_: deepcopy(plan), tmp_path/'original')
    bundle = deepcopy(original['acquisition_snapshot'])
    bundle['case']['info']['regularMarketPrice'] = 81.
    if receipt == 'missing':
        bundle['case']['sources']['profile'].pop('retrieved_at', None)
    elif receipt == 'future':
        bundle['case']['sources']['profile']['retrieved_at'] = '2099-09-28T12:00:00+00:00'
    from bellomberg.valuation.sector_analysis import _hash
    bundle['snapshot_id'] = _hash({key: value for key, value in bundle.items() if key != 'snapshot_id'})
    before = deepcopy(bundle)
    refreshed = generate_valuation(bundle['case']['ticker'], output_dir=str(tmp_path/'refreshed'), prepared_bundle=bundle)
    assert candidate_model_usability(refreshed)['usable']
    quote = refreshed.get('market_quote') or {}
    assert quote.get('status') == ('ok' if receipt == 'valid' else 'source_unavailable')
    assert quote['intrinsic_value_applicable'] is False
    assert quote['fv_basis'] == 'not_applicable_observed_exposure'
    assert refreshed['price'] == original['price'] == 11.
    assert refreshed['valuation_date'] == original['valuation_date'] == '2026-06-30'
    assert refreshed['exposure_analysis'] == original['exposure_analysis']
    assert bundle == before
    assert all(refreshed.get('fair_value_'+scenario) is None and quote.get('upside_'+scenario+'_pct') is None
               for scenario in ('bear', 'base', 'bull'))
    if receipt == 'valid':
        assert quote['price'] == 81. and quote['price_model'] == 11.
        assert quote['price_model_as_of'] == '2026-09-25'
        assert quote['retrieved_at'] == bundle['case']['sources']['profile']['retrieved_at']


def test_commodity_public_projection_rebuilds_same_observations_from_exact_original_sources(tmp_path):
    from bellomberg.valuation.trade_idea_model import prepare, public_model_payload, prepare_shareable_model, model_exhibits
    q, plan = primary_sourced(tmp_path)
    plan['analysis_rationale'] = 'PRIVATE_RATIONALE_CANARY'
    payload = prepare(q, lambda *_: deepcopy(plan), tmp_path/'private')
    payload['private_note'] = 'PRIVATE_NOTE_CANARY'
    projected = public_model_payload(payload)
    assert projected['status'] == 'ready', projected['reasons']
    encoded = json.dumps(projected, sort_keys=True)
    assert 'PRIVATE_NOTE_CANARY' not in encoded and 'PRIVATE_RATIONALE_CANARY' not in encoded
    assert str(tmp_path) not in encoded and 'archive_path' not in encoded
    public = prepare_shareable_model(payload, tmp_path/'public')
    assert public['snapshot_id'] != payload['snapshot_id'] and public['generation_id'] != payload['generation_id']
    for key in ('accounting', 'fees', 'liquidity', 'quotation', 'accounting_nav_per_share', 'gold_nav_ratio', 'liability_nav_ratio'):
        assert public['exposure_analysis'][key] == payload['exposure_analysis'][key]
    assert model_exhibits(public)['status'] == 'complete'
    assert public['exposure_analysis']['investment_leverage'] is None
    assert all(public.get('fair_value_'+scenario) is None for scenario in ('bear','base','bull'))


def test_commodity_public_projection_rejects_resealed_normalized_observation_and_missing_original(tmp_path):
    from bellomberg.valuation.trade_idea_model import prepare, public_model_payload
    q, plan = primary_sourced(tmp_path)
    payload = prepare(q, lambda *_: deepcopy(plan), tmp_path/'private')
    for fault in ('normalized', 'original', 'record_proof'):
        changed = deepcopy(payload)
        documents = changed['preparation']['review_basis']['dossier']['documents']
        if fault == 'original':
            documents[:] = [row for row in documents if (row.get('metadata') or {}).get('normalizer')]
        elif fault == 'normalized':
            doc = next(row for row in documents if (row.get('metadata') or {}).get('normalizer'))
            body = json.loads(doc['text']); body['observations'][0]['value']['name'] = 'FORGED_PUBLIC_TEXT'
            doc['text'] = json.dumps(body, sort_keys=True); doc['sha256'] = sha256(doc['text'].encode()).hexdigest()
        else:
            changed['preparation']['proposal']['plan']['model']['risks']['record_pointer'] = '/observations/0'
        assert public_model_payload(changed)['status'] == 'blocked', fault


@pytest.mark.parametrize('fault', ['nav_per_share', 'squared_weight'])
def test_commodity_finite_observations_do_not_publish_overflowing_derived_metrics(fault):
    from bellomberg.valuation.trade_idea_commodity_exposure import project
    model = values()
    if fault == 'nav_per_share':
        model['holdings'] = {'assets': {'gold': 1e308}, 'liabilities': {'sponsor_fee_payable': 0.},
            'reported': {'total_assets': 1e308, 'total_liabilities': 0., 'net_assets': 1e308, 'shares': 1e-308},
            'as_of': '2026-06-30', 'currency': 'USD'}
    else:
        model['holdings'] = {'assets': {'gold': 1e-150}, 'liabilities': {'sponsor_fee_payable': 0.},
            'reported': {'total_assets': 1e-150, 'total_liabilities': 0., 'net_assets': 1e-308, 'shares': 1.},
            'as_of': '2026-06-30', 'currency': 'USD'}
    result = project(model, ticker='SYNTH-GOLD', instrument='etf', cutoff='2026-09-28')
    assert result['status'] == 'incomplete' and any('representable' in reason for reason in result['issues'])


def test_nonzero_observed_ratios_are_readable_as_fractions_and_volume_is_integral(tmp_path):
    from bellomberg.valuation.trade_idea_model import prepare, model_exhibits
    from openpyxl import load_workbook
    from pathlib import Path
    q, plan = sourced(tmp_path)
    payload = prepare(q, lambda *_: deepcopy(plan), tmp_path/'model')
    workbook = load_workbook(payload['path'])
    for row in (12,13,14,17,18):
        assert workbook['Summary'].cell(row,4).number_format == '0.000000'
        assert workbook['Summary'].cell(row,3).value == 'ratio'
    assert workbook['Summary']['D19'].number_format == '#,##0'
    refs = payload['calculation_details']['workbook_bindings']['input_refs']
    for driver, key in (('fees','annualized_expense_ratio'),('liquidity','spread_ratio')):
        ref = next(row for row in refs if row['driver'] == driver and row['path'] == [key])
        cell = ref['cell'].split('!')[1]
        assert workbook['Observed inputs'][cell].number_format == '0.000000'
    workbook['Summary']['D18'].number_format = '#,##0.00'
    workbook.save(payload['path']); workbook.close()
    payload['workbook_sha256'] = sha256(Path(payload['path']).read_bytes()).hexdigest()
    assert model_exhibits(payload)['status'] == 'blocked'


@pytest.mark.parametrize('suffix', ['?pm_note=PRIVATE_SOURCE_QUERY_CANARY', '?PM_NOTE=PRIVATE_SOURCE_QUERY_CANARY',
    '#PRIVATE_SOURCE_QUERY_CANARY', '%3Fpm_note=PRIVATE_SOURCE_QUERY_CANARY',
    '%253fpm_note=PRIVATE_SOURCE_QUERY_CANARY', '%23PRIVATE_SOURCE_QUERY_CANARY'])
def test_public_projection_does_not_publish_unproven_source_url_parameters(tmp_path,monkeypatch,suffix):
    import test_commodity_trust_evidence as originals
    from bellomberg.valuation.trade_idea_model import prepare, public_model_payload
    source = originals.source
    def private_url(path,url,text):
        document=source(path,url,text)
        if url.endswith('/product'): document['url']+=suffix
        return document
    monkeypatch.setattr(originals,'source',private_url)
    q,plan=primary_sourced(tmp_path)
    payload=prepare(q,lambda *_:deepcopy(plan),tmp_path/'private')
    projected=public_model_payload(payload)
    assert projected['status']=='blocked' and projected['bundle'] is None
    assert projected['sources']==[] and 'PRIVATE_SOURCE_QUERY_CANARY' not in json.dumps(projected)


def test_observation_workbook_has_room_for_full_dollars_dates_and_risk_text(tmp_path):
    from bellomberg.valuation.trade_idea_model import prepare
    from openpyxl import load_workbook
    q,plan=sourced(tmp_path)
    payload=prepare(q,lambda *_:deepcopy(plan),tmp_path/'model')
    workbook=load_workbook(payload['path'])
    assert workbook['Summary'].column_dimensions['D'].width >= 28
    assert workbook['Observed inputs'].column_dimensions['D'].width >= 56
    assert workbook['Summary']['D24'].alignment.wrap_text
    record=next(row for row in workbook['Observed inputs'] if row[1].value=='risks')
    assert record[3].alignment.wrap_text
    workbook.close()
