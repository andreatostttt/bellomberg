"""Declared synthetic arithmetic/adversarial tests; live WDP proof is archived separately."""
from copy import deepcopy
import pytest
from openpyxl import load_workbook
from bellomberg.valuation.property_nav_requirements import REPORTED_POLICY, REPORTED_SCHEMA, reported_policy_selected
from bellomberg.valuation.property_reported_nav import ASSETS, LIABILITIES, SHOCKS, TARGET_BASIS, project_reported
from bellomberg.valuation.method_registry import get_method_requirements, requirements_for_records

DAY='2026-09-28'; ON='2026-06-30'


def synthetic_model():
    assets={k:0. for k in ASSETS};assets.update(investment_property=100.,cash=10.,participations=20.)
    liabilities={k:0. for k in LIABILITIES};liabilities['noncurrent_debt']=40.
    return {'perimeter':{'entity':'SYNTH-REPORTED-PROP','currency':'EUR','share_class':'synthetic dividend shares'},
        'calendar':{'valuation_date':ON,'periods':[],'discount_convention':'snapshot'},
        'quotation':{'financial_currency':'EUR','quote_currency':'EUR','quote_unit':'EUR','quote_units_per_currency':1,
            'financial_to_quote_rate':1,'shares_per_quote':1,'share_class':'synthetic dividend shares','price':10.,'price_as_of':ON},
        'shares':10.,'policy':deepcopy(REPORTED_POLICY),
        'balance_sheet':{'assets':assets,'liabilities':liabilities,'reported_assets':130.,'reported_liabilities':40.,
            'reported_parent_equity':90.,'reported_minority_equity':0.,'money_precision':0.001},
        'property_reconciliation':{'standing':80.,'development':15.,'land':5.,'held_for_sale':0.,
            'jv_property_memorandum':25.,'portfolio_fair_value':125.,'leasehold_rights_in_standing':8.,
            'jv_ownership':{'synthetic JV':'50% documented synthetic equity method'},'jv_basis':'equity_method_property_memorandum_only'},
        'epra_bridge':{'basis':'EPRA_NTA','adjustments':{'deferred_tax_property':4.,'financial_instruments':-2.,'intangibles':0.},
            'reported_nav':92.,'reported_per_share':9.2,'per_share_decimals':1},
        'claims':{'entitled_shares_m':10.,'epra_diluted_shares_m':10.,'reported_diluted_ifrs_nav':90.,'hybrid_claims':0.,
            'dilution_basis':'reported_EPRA_after_options_convertibles_other_equity','post_balance_events':'Explicit synthetic no subsequent valuation rollforward.'},
        'restrictions':{'covenants':{name:{'observed':v,'threshold':limit,'operator':operator,'basis':'Explicit synthetic covenant convention'}
            for name,v,limit,operator in [('interest_coverage',4.,1.5,'>='),('statutory_gearing',.3,.65,'<'),
                ('consolidated_gearing',.31,.65,'<'),('non_prelet_development',.01,.15,'<='),('subsidiary_debt',.05,.3,'<=')]},
            'guarantees':{'synthetic_parent_guarantee':3.},'encumbrance':'Explicit synthetic negative pledge and conditional repayment.', 'compliance_date':ON},
        'development':{'pipeline_total':30.,'invested':10.,'cost_to_come':20.,'undrawn_confirmed_lines':50.,
            'debt_maturities_to_end_next_year':15.,'funding_basis':'reported_aggregate_not_project_financing_or_future_profit',
            'pipeline_scope':'reported_projects_excluding_energy_and_land_reserves'}}


def synthetic_scenarios():
    return {s:{'asset_shocks':{k:(overrun if k=='development_cost_overrun' else shock) for k in SHOCKS},
        'nav_target':1.,'target_basis':TARGET_BASIS} for s,shock,overrun in [('bear',-.1,.05),('base',0.,0.),('bull',.1,0.)]}


def synthetic_bundle(model=None,scenarios=None):
    from bellomberg.valuation.sector_analysis import prepare_sector_analysis
    model=synthetic_model() if model is None else model;scenarios=synthetic_scenarios() if scenarios is None else scenarios
    records=[]
    for scope,values in [('model',model),*scenarios.items()]:
        for driver,value in values.items():
            field,unit,basis,timing,shape,owner=REPORTED_SCHEMA[driver]
            records.append({'driver':driver,'field':field,'value':deepcopy(value),'scenario':scope,'entity':'SYNTH-REPORTED-PROP',
                'period':ON,'unit':unit,'accounting_basis':basis,'source_id':'https://example.org/declared-synthetic/property',
                'as_of':DAY,'valid_until':DAY,'kind':'historical' if scope=='model' and driver not in ('perimeter','calendar','policy') else 'analyst_estimate',
                'rationale':'Declared synthetic arithmetic case; no live issuer qualification.'})
    return prepare_sector_analysis('SYNTH-REPORTED-PROP',as_of=DAY,providers={
        'profile':lambda *a,**k:{'status':'ok','source_id':'synthetic','as_of':DAY,'data':{'evidence':[
            {'field':f,'value':v,'source_id':'synthetic','as_of':DAY} for f,v in [('instrument','equity'),('business_model','property_owner')]]}},
        'method_inputs':lambda *a,**k:{'status':'ok','source_id':'synthetic','as_of':DAY,'records':records}},
        user_context={'analysis_context':{'scenario_rationale':{s:'Explicit synthetic value sensitivity, not a source forecast.' for s in scenarios}}})


def test_exact_opt_in_preserves_legacy_requirements():
    legacy=get_method_requirements('property_nav')
    assert requirements_for_records('property_nav',[])==legacy
    record={'driver':'policy','scenario':'model','value':deepcopy(REPORTED_POLICY)}
    assert 'ffo_affo_bridge' not in {f['field'] for f in requirements_for_records('property_nav',[record])['fields']}
    for change in ({'version':'2'},{'extra':'unconsumed'}):
        bad=deepcopy(record);bad['value'].update(change)
        assert requirements_for_records('property_nav',[bad])==legacy
    assert requirements_for_records('property_nav',[record,record])==legacy
    assert not reported_policy_selected('malformed')


def test_asset_nav_independent_oracle_excludes_second_jv_and_committed_capex():
    issues=[];results=project_reported(synthetic_model(),synthetic_scenarios(),lambda d,r:issues.append((d,r)))
    assert not issues
    assert {s:r['fair_value_per_share'] for s,r in results.items()}==pytest.approx({'bear':7.9,'base':9.2,'bull':10.4})
    assert results['base']['epra_nta_reconstructed']==92
    assert results['base']['jv_property_memorandum_excluded_from_nav']==25
    assert results['base']['development_cost_to_come_memorandum']==20
    assert results['bear']['development_overrun_deduction']==1


def test_historical_gate_does_not_require_or_create_scenario_assumptions():
    issues=[];result=project_reported(synthetic_model(),{},lambda d,r:issues.append((d,r)),historical_only=True)
    assert not issues
    assert result['epra_nta_reconstructed']==92
    assert set(result)=={'reconciliations','ifrs_parent_equity_reconstructed','epra_nta_reconstructed'}
    assert 'fair_value_per_share' not in result


@pytest.mark.parametrize('fault',['asset_missing','liability_missing','liability_double','jv_double','development_double',
    'nta_wrong','dilution','hybrid','missing_covenant','breached_covenant','future_compliance','opaque_precision',
    'leasehold_overclaim','extra_claim','bad_pipeline','extra_shock','negative_overrun','wrong_target'])
def test_incomplete_or_false_reconciliations_block_common_generation(tmp_path,fault):
    from bellomberg.valuation.dcf_engine import generate_valuation
    m=synthetic_model();sc=synthetic_scenarios()
    if fault=='asset_missing':del m['balance_sheet']['assets']['cash']
    if fault=='liability_missing':del m['balance_sheet']['liabilities']['current_accruals']
    if fault=='liability_double':m['balance_sheet']['liabilities']['noncurrent_debt']*=2
    if fault=='jv_double':m['balance_sheet']['assets']['participations']+=25
    if fault=='development_double':m['property_reconciliation']['development']+=10
    if fault=='nta_wrong':m['epra_bridge']['reported_nav']+=1
    if fault=='dilution':m['claims']['epra_diluted_shares_m']=12
    if fault=='hybrid':m['claims']['hybrid_claims']=1
    if fault=='missing_covenant':del m['restrictions']['covenants']['subsidiary_debt']
    if fault=='breached_covenant':m['restrictions']['covenants']['interest_coverage']['observed']=1.
    if fault=='future_compliance':m['restrictions']['compliance_date']='2026-12-31'
    if fault=='opaque_precision':m['balance_sheet']['money_precision']=1000000
    if fault=='leasehold_overclaim':m['property_reconciliation']['leasehold_rights_in_standing']=100
    if fault=='extra_claim':m['claims']['unproved_no_options']=True
    if fault=='bad_pipeline':m['development']['cost_to_come']=100
    if fault=='extra_shock':sc['base']['asset_shocks']['ignored']=0
    if fault=='negative_overrun':sc['base']['asset_shocks']['development_cost_overrun']=-.1
    if fault=='wrong_target':sc['base']['target_basis']='EPRA_NRV_called_bull'
    result=generate_valuation('SYNTH-REPORTED-PROP',prepared_bundle=synthetic_bundle(m,sc),output_dir=str(tmp_path))
    assert not result['valuation_usability']['usable']
    assert result.get('fair_value_base') is None
    wb=load_workbook(result['path']);assert 'Reported NAV base' not in wb;wb.close()


def test_common_workbook_has_live_equations_and_source_bound_residuals(tmp_path):
    from bellomberg.valuation.dcf_engine import generate_valuation
    result=generate_valuation('SYNTH-REPORTED-PROP',prepared_bundle=synthetic_bundle(),output_dir=str(tmp_path))
    assert result['valuation_usability']['usable'],result.get('error')
    wb=load_workbook(result['path'])
    assert wb['Model Checks']['C9'].value==pytest.approx(9.2)
    assert wb['Reported NAV base']['D11'].data_type=='f'
    assert wb['Reported NAV base']['D24'].data_type=='f'
    assert wb['Summary']['E8'].data_type=='f'
    assert all(len(c.value)<8192 for ws in wb for row in ws for c in row if c.data_type=='f')
    text=' '.join(str(c.value) for ws in wb for row in ws for c in row)
    assert 'future compliance' in text or 'compliance futura' in text
    assert 'ffo_bridge' not in result['input_consumption']['consumed_fields']
    wb.close()


def test_reported_rounding_residual_is_visible_and_bounded():
    model=synthetic_model();model['balance_sheet']['reported_parent_equity']+=.001
    model['claims']['reported_diluted_ifrs_nav']=model['balance_sheet']['reported_parent_equity']
    issues=[];r=project_reported(model,synthetic_scenarios(),lambda d,t:issues.append((d,t)))
    assert not issues
    check=r['base']['reconciliations']['parent_equity']
    assert check['residual']==pytest.approx(-.001)
    assert check['rounding_bound']==pytest.approx(.0115)
    assert r['base']['epra_nta_reconstructed']==92
