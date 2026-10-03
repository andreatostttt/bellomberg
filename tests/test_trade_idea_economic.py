"""Trade Idea source gate and exact common-model vertical, entirely offline."""
from copy import deepcopy
from pathlib import Path

import pytest

from bellomberg.valuation.trade_idea_model import qualification, prepare, source_fingerprint, model_exhibits
from test_input_preparation import _documents as _legacy_documents, _operating_plan as _legacy_operating_plan
from test_sector_analysis import providers_for as _legacy_providers_for, DAY


IDENTITY = {'ticker': 'SYNTH-EXT', 'name': 'Synthetic issuer', 'exchange': 'TEST',
            'currency': 'EUR', 'status': 'confirmed'}


def providers_for(*args, **kwargs):
    """Explicit synthetic public quote venue; common legacy fixtures are untouched."""
    providers = _legacy_providers_for(*args, **kwargs)
    original = providers['profile']
    def profile(*profile_args, **profile_kwargs):
        result = original(*profile_args, **profile_kwargs)
        result['as_of'] = profile_kwargs.get('as_of', DAY)
        result['retrieved_at'] = result['as_of']+'T12:00:00+00:00'
        result['data']['info'].update(synthetic_observed_quote_info(
            profile_args[0], profile_kwargs.get('as_of', DAY)))
        return result
    return {**providers, 'profile': profile}


def synthetic_observed_quote_info(ticker, as_of=DAY, *, currency='EUR', price=10., exchange='TEST'):
    """Explicit fictional observed quote; no historical model price is reused."""
    from datetime import datetime, timezone
    stamp = datetime.fromisoformat(as_of+'T12:00:00').replace(tzinfo=timezone.utc).timestamp()
    return {'symbol': ticker, 'regularMarketPrice': price, 'regularMarketTime': stamp,
            'exchangeTimezoneName': 'UTC', 'fullExchangeName': exchange, 'exchange': exchange,
            'currency': currency, 'quoteSourceName': 'Fictional transport observation'}


@pytest.mark.parametrize('fault', ['old_price_acquired_today', 'future_price', 'missing_time',
    'missing_timezone', 'boolean_price', 'nonfinite_price', 'nan_price', 'missing_source',
    'old_acquisition', 'old_source_and_price', 'future_acquisition', 'missing_currency', 'source_unavailable'])
def test_free_gate_requires_a_current_dated_quote_separate_from_the_model_snapshot(tmp_path, fault):
    from datetime import datetime, timezone
    providers = providers_for()
    original = providers['profile']
    def profile(*args, **kwargs):
        result = original(*args, **kwargs)
        info = result['data']['info']
        if fault in ('old_price_acquired_today', 'future_price'):
            day = '2026-06-30' if fault.startswith('old') else '2026-09-11'
            info['regularMarketTime'] = datetime.fromisoformat(day+'T12:00:00').replace(tzinfo=timezone.utc).timestamp()
        elif fault == 'missing_time': info.pop('regularMarketTime')
        elif fault == 'missing_timezone': info.pop('exchangeTimezoneName')
        elif fault == 'boolean_price': info['regularMarketPrice'] = True
        elif fault == 'nonfinite_price': info['regularMarketPrice'] = float('inf')
        elif fault == 'nan_price': info['regularMarketPrice'] = float('nan')
        elif fault == 'missing_source': result['source_id'] = None
        elif fault == 'old_source_and_price':
            result['as_of'] = '2026-06-30'
            info['regularMarketTime'] = datetime.fromisoformat('2026-06-30T12:00:00').replace(tzinfo=timezone.utc).timestamp()
        elif fault == 'future_acquisition': result['as_of'] = '2026-09-11'
        elif fault == 'missing_currency': info.pop('currency')
        elif fault == 'source_unavailable': result['status'] = 'source_error'
        else: result['as_of'] = '2026-06-30'
        return result
    providers['profile'] = profile
    result = qualification('SYNTH-EXT', IDENTITY, DAY, archive_root=tmp_path,
        providers=providers, source_report={'documents': _documents(), 'source_plan': _operating_plan()})
    assert result['status'] == 'blocked', result
    calls = []
    with pytest.raises(ValueError, match='Fonti non qualificate'):
        prepare(result, lambda *_: calls.append('forbidden-paid'), tmp_path/'model')
    assert calls == []


def test_current_price_receipt_never_qualifies_fair_value_or_uses_the_model_price(tmp_path):
    from bellomberg.valuation.trade_idea_model import _current_quotation_evidence
    result = qualified(tmp_path)
    bundle = deepcopy(result['bundle'])
    bundle['case']['info']['regularMarketPrice'] = 20.
    plan = deepcopy(result['source_report']['source_plan'])
    problem, basis = _current_quotation_evidence(bundle, plan, DAY)
    assert problem is None and basis['price'] == 20.
    assert basis['comparison_qualified'] is False
    assert basis['comparison_status'] == 'not_assessed_before_model'
    original_price = plan['model']['quotation']['value']['price']
    plan['model']['quotation']['value']['financial_currency'] = 'USD'
    problem, basis = _current_quotation_evidence(bundle, plan, DAY)
    assert problem is None and basis['price'] == 20.
    assert basis['comparison_qualified'] is False and basis['comparison_status'] == 'fx_not_rolled'
    assert plan['model']['quotation']['value']['price'] == original_price


@pytest.mark.parametrize('fault', ['different_exchange', 'missing_exchange', 'different_full_name'])
def test_confirmed_quote_venue_requires_an_exact_public_profile_match(tmp_path, fault):
    providers = providers_for()
    original = providers['profile']
    def profile(*args, **kwargs):
        result = original(*args, **kwargs)
        info = result['data']['info']
        if fault == 'missing_exchange':
            info.pop('exchange'); info.pop('fullExchangeName')
        elif fault == 'different_full_name':
            info['fullExchangeName'] = 'OTHER MARKET'
        else:
            info.pop('fullExchangeName'); info['exchange'] = 'OTHER MARKET'
        return result
    providers['profile'] = profile
    calls = []
    result = qualification('SYNTH-EXT', IDENTITY, DAY, archive_root=tmp_path,
        providers=providers, collector=lambda *_a, **_k: calls.append('collector'))
    assert result['status'] == 'blocked'
    assert any('Borsa del profilo' in reason for reason in result['reasons'])
    assert calls == []
    with pytest.raises(ValueError, match='Fonti non qualificate'):
        prepare(result, lambda *_: pytest.fail('paid callback after quote venue mismatch'), tmp_path/'model')


def _documents(**kwargs):
    """Synthetic source contains both opening claims; zero is observed here."""
    import json
    from hashlib import sha256
    observations = [{'field': 'enterprise_equity_bridge', 'driver': driver, 'value': 0.,
        'entity': 'SYNTH-GROUP', 'period': '2025-12-31', 'unit': 'EUR million',
        'accounting_basis': 'valuation'} for driver in ('net_debt', 'equity_adjustments')]
    text = json.dumps({'synthetic': True, 'observations': observations}, sort_keys=True)
    return [*_legacy_documents(**kwargs), {'id': 'synthetic-opening-claims',
        'url': 'https://example.org/issuer/synthetic-opening-claims', 'published_at': '2026-09-09',
        'text': text, 'sha256': sha256(text.encode()).hexdigest()}]


def _operating_plan(years=10):
    plan = _legacy_operating_plan(years)
    for drivers in plan['scenarios'].values():
        for index, driver in enumerate(('net_debt', 'equity_adjustments')):
            drivers[driver].update(kind='historical', evidence_ids=['synthetic-opening-claims'],
                record_pointer='/observations/'+str(index),
                rationale='Explicit synthetic source observation at the opening economic date')
    return plan


def qualified(tmp_path, *, plan=None, documents=None):
    return qualification('SYNTH-EXT', IDENTITY, DAY, archive_root=tmp_path,
        providers=providers_for(), source_report={'documents': documents or _documents(),
        'source_plan': deepcopy(plan or _operating_plan())})


def test_confirmed_foreign_quote_cannot_be_replaced_by_primary_report_currency(tmp_path):
    providers = providers_for()
    original = providers['profile']
    def profile(*args, **kwargs):
        value = original(*args, **kwargs)
        value['data']['info']['currency'] = 'USD'
        value['data']['info']['financialCurrency'] = 'EUR'
        return value
    providers['profile'] = profile
    result = qualification('SYNTH-EXT', {**IDENTITY, 'currency': 'USD'}, DAY,
        archive_root=tmp_path, providers=providers,
        source_report={'documents': _documents(), 'source_plan': _operating_plan()})
    assert result['status'] == 'blocked'
    assert any('confirmed instrument' in reason for reason in result['reasons'])
    with pytest.raises(ValueError, match='Fonti non qualificate'):
        prepare(result, lambda *_: pytest.fail('paid callback after quotation identity mismatch'), tmp_path/'model')


@pytest.mark.parametrize('accepted,financial,currency,unit', [
    ('EUR', 'USD', 'EUR', 'EUR'), ('GBX', 'USD', 'GBP', 'GBX'),
    ('GBp', 'USD', 'GBP', 'GBX'),
])
def test_quote_identity_keeps_financial_currency_separate(accepted, financial, currency, unit):
    from bellomberg.valuation.trade_idea_model import _quotation_identity_problem
    plan = {'model': {'quotation': {'value': {'financial_currency': financial,
        'quote_currency': currency, 'quote_unit': unit}}}}
    assert _quotation_identity_problem({'currency': accepted}, plan) is None


@pytest.mark.parametrize('driver', ['net_debt', 'equity_adjustments'])
@pytest.mark.parametrize('fault', ['estimate_zero', 'missing', 'unproved', 'wrong_period', 'scenario_conflict'])
def test_missing_opening_claims_never_become_analyst_zeros_before_paid(tmp_path, driver, fault):
    plan, documents = _operating_plan(), _documents()
    item = plan['scenarios']['base'][driver]
    if fault == 'estimate_zero':
        item.update(kind='analyst_estimate')
        item.pop('record_pointer')
    elif fault == 'missing':
        plan['scenarios']['base'].pop(driver)
    elif fault == 'unproved':
        item['record_pointer'] = '/observations/99'
    elif fault == 'scenario_conflict':
        item['value'] = 1.
    else:
        import json
        from hashlib import sha256
        doc = documents[-1]
        body = json.loads(doc['text'])
        next(row for row in body['observations'] if row['driver'] == driver)['period'] = '2026-01-01'
        doc['text'] = json.dumps(body, sort_keys=True)
        doc['sha256'] = sha256(doc['text'].encode()).hexdigest()
    q = qualified(tmp_path, plan=plan, documents=documents)
    assert q['status'] == 'blocked', q
    assert any(driver in reason for reason in q['reasons'])
    with pytest.raises(ValueError, match='Fonti non qualificate'):
        prepare(q, lambda *_: pytest.fail('paid proposer after unproved opening claim'), tmp_path)


@pytest.mark.parametrize('driver', ['net_debt', 'equity_adjustments'])
@pytest.mark.parametrize('fault', ['estimate_zero', 'changed_balance', 'missing'])
def test_preparer_cannot_replace_qualified_opening_facts_before_desks(tmp_path, driver, fault):
    q = qualified(tmp_path)
    assert q['status'] == 'qualified'
    plan = _operating_plan()
    item = plan['scenarios']['base'][driver]
    if fault == 'estimate_zero':
        item.update(kind='analyst_estimate')
        item.pop('record_pointer')
    elif fault == 'changed_balance':
        item['value'] += 1
    else:
        plan['scenarios']['base'].pop(driver)
    calls = []
    def proposer(*_args):
        calls.append('budgeted callback')
        return deepcopy(plan)
    payload = prepare(q, proposer, tmp_path / 'model')
    assert len(calls) == 1
    assert payload['valuation_usability']['usable'] is False
    assert payload['preparation']['status'] == 'incomplete'
    assert not payload.get('path')
    assert driver in payload['error']


@pytest.mark.parametrize('driver', ['shares', 'quotation'])
def test_impossible_source_denominators_block_before_paid(tmp_path, driver):
    plan = _operating_plan()
    if driver == 'shares':
        plan['model'][driver]['value'] = -10
    else:
        plan['model'][driver]['value']['price'] = 0
    q = qualified(tmp_path, plan=plan)
    assert q['status'] == 'blocked'
    assert any('positive finite' in reason for reason in q['reasons'])


def test_free_source_plan_proof_then_exact_common_workbook(tmp_path):
    q = qualified(tmp_path)
    assert q['status'] == 'qualified', q['reasons']
    payload = prepare(q, lambda dossier, contract: deepcopy(_operating_plan()), tmp_path / 'model')
    assert payload['valuation_usability']['usable'], payload.get('error')
    assert payload['preparation']['status'] == 'prepared'
    assert payload['acquisition_snapshot']['snapshot_id'] == payload['snapshot_id']
    assert payload['generation_id'] and payload['workbook_sha256']
    assert Path(payload['path']).is_file()
    from openpyxl import load_workbook
    workbook = load_workbook(payload['path'], data_only=False)
    assert workbook['Summary']['E8'].data_type == 'f'
    assert workbook['base']['D26'].data_type == 'f'
    workbook.close()
    exhibits = model_exhibits(payload)
    assert exhibits['status'] == 'complete', exhibits['reasons']
    ids = {row['id'] for row in exhibits['exhibits']}
    assert {'scenarios', 'drivers', 'valuation_bridge', 'sensitivity'} <= ids
    scenario = next(row for row in exhibits['exhibits'] if row['id'] == 'scenarios')
    assert scenario['rows'][1][1] == payload['fair_value_base']
    assert scenario['cell_refs'][1][1] == "'Summary'!E8"
    sidecar = __import__('json').loads(Path(payload['path']).with_suffix('.payload.json').read_text(encoding='utf-8'))
    sidecar_exhibits = model_exhibits(sidecar, workbook_path=payload['path'])
    assert sidecar_exhibits == exhibits


def test_missing_or_fabricated_opening_proof_blocks_before_proposer(tmp_path):
    plan = _operating_plan()
    plan['model']['historical_revenue'].update(value=999., quoted_value=999.)
    q = qualified(tmp_path, plan=plan)
    assert q['status'] == 'blocked'
    assert any('historical_revenue' in reason for reason in q['reasons'])
    with pytest.raises(ValueError, match='Fonti non qualificate'):
        prepare(q, lambda *_: pytest.fail('paid proposer after missing opening proof'), tmp_path)


def test_document_url_count_never_qualifies_sources(tmp_path):
    q = qualification('SYNTH-EXT', IDENTITY, DAY, archive_root=tmp_path,
                      providers=providers_for(), source_report={'documents': _documents(), 'preparation_ready': True})
    assert q['status'] == 'blocked'
    assert any('quotazione' in reason.lower() for reason in q['reasons'])


def test_collector_readiness_does_not_certify_the_economic_bridge(tmp_path):
    documents = _documents()
    documents[0]['metadata'] = {'report_date': DAY}
    documents.append(dict(deepcopy(documents[0]), id='price-SYNTH-EXT-' + DAY))
    q = qualification('SYNTH-EXT', IDENTITY, DAY, archive_root=tmp_path,
        providers=providers_for(), source_report={'documents': documents, 'preparation_ready': True,
        'selection': {'opening_date': DAY}, 'balance_sheet': {'status': 'ready'}, 'source_gaps': []})
    assert q['status'] == 'blocked'
    assert any('riconciliazione economica' in reason for reason in q['reasons'])


def test_retrieval_time_and_archive_path_do_not_change_semantic_pin(tmp_path):
    q = qualified(tmp_path)
    copy = deepcopy(q)
    copy['source_report']['documents'].reverse()
    copy['source_report']['documents'][0]['metadata'] = {'archive_path': 'elsewhere', 'downloaded_at': 'tomorrow'}
    # An absent metadata object and an empty object describe the same evidence.
    next(doc for doc in q['source_report']['documents']
         if doc['id'] == copy['source_report']['documents'][0]['id'])['metadata'] = {}
    assert source_fingerprint(copy) == source_fingerprint(q)


def test_mutated_pinned_source_rejected_before_proposer(tmp_path):
    q = qualified(tmp_path)
    q['source_report']['documents'][0]['sha256'] = '0' * 64
    with pytest.raises(ValueError, match='modificate'):
        prepare(q, lambda *_: pytest.fail('paid proposer after source mutation'), tmp_path)


def test_all_existing_intrinsic_adapters_have_preparation_schema():
    from bellomberg.valuation.preparation_methods import supported_methods, method_schema
    methods = supported_methods()
    assert len(methods) == 14
    for method in methods:
        schema, entities = method_schema(method)
        assert schema and 'perimeter' in schema and 'quotation' in schema


def test_paid_revision_receives_exact_before_and_consumes_requested_driver(tmp_path):
    from bellomberg.valuation.trade_idea_model import revise
    q = qualified(tmp_path)
    before = prepare(q, lambda *_: deepcopy(_operating_plan()), tmp_path / 'before')
    records = [deepcopy(row) for row in before['acquisition_snapshot']['case']['records'] if row['driver'] == 'wacc']
    for record in records:
        record['value'] = .12
    changes = {'method_records': records, 'rationale': 'Explicit higher discount-rate challenge',
               'evidence_ids': [q['source_report']['documents'][0]['id']]}
    calls = []
    def propose(dossier, _contract):
        calls.append(deepcopy(dossier))
        plan = deepcopy(_operating_plan())
        for drivers in plan['scenarios'].values():
            drivers['wacc']['value'] = .12
        return plan
    after = revise(before, changes, qualification=q, propose=propose, output_dir=tmp_path / 'after')
    assert after['valuation_usability']['usable'], after.get('error')
    context = calls[0]['preparation_context']
    assert context['before']['generation_id'] == before['generation_id']
    assert context['reviewer_changes'] == changes
    assert after['fair_value_base'] != before['fair_value_base']
    lineage = after['preparation']['provenance']['trade_idea_revision']
    assert lineage['previous_generation_id'] == before['generation_id']
    assert lineage['rationale'] == changes['rationale']
    assert after['preparation']['review_basis']['dossier']['preparation_context'] == context
    with pytest.raises(ValueError, match='non applicato'):
        revise(before, changes, qualification=q, propose=lambda *_: deepcopy(_operating_plan()),
               output_dir=tmp_path / 'ignored')
    q['source_report']['documents'][0]['sha256'] = '0' * 64
    with pytest.raises(ValueError, match='modificate'):
        revise(before, changes, qualification=q, propose=lambda *_: pytest.fail('paid after drift'),
               output_dir=tmp_path / 'drift')


def test_public_model_is_regenerated_without_personal_text_or_workbook_parts(tmp_path):
    from bellomberg.valuation.trade_idea_model import public_model_payload, prepare_shareable_model
    import json
    from zipfile import ZipFile
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    q = qualified(tmp_path)
    payload = prepare(q, lambda *_: deepcopy(_operating_plan()), tmp_path / 'private')
    canary = 'PERSONAL_PM_CANARY_927415'
    payload['preparation']['provenance']['notes'] = canary
    payload['acquisition_snapshot']['case']['sources']['profile']['data']['info']['personal_note'] = canary
    # Deliberately rehash after injecting a valid but private source extension.
    from bellomberg.valuation.sector_analysis import _hash
    bundle = payload['acquisition_snapshot']
    bundle['snapshot_id'] = _hash({key: value for key, value in bundle.items() if key != 'snapshot_id'})
    workbook = load_workbook(payload['path'])
    workbook.properties.creator = canary
    workbook['Summary']['A1'] = canary
    workbook['Summary']['E8'].comment = Comment(canary, canary)
    workbook.create_sheet(canary).sheet_state = 'hidden'
    workbook.save(payload['path']); workbook.close()
    projected = public_model_payload(payload)
    assert projected['status'] == 'ready', projected['reasons']
    assert canary not in json.dumps(projected)
    shared = prepare_shareable_model(payload, tmp_path / 'public')
    assert shared['snapshot_id'] != payload['snapshot_id']
    assert shared['generation_id'] != payload['generation_id']
    assert shared['workbook_sha256'] != payload['workbook_sha256']
    assert shared['fair_value_base'] == payload['fair_value_base']
    with ZipFile(shared['path']) as package:
        assert all(canary.encode() not in package.read(name) for name in package.namelist())
