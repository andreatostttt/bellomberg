"""Explicit commodity-trust accounting and separately dated market observations.

The contract is observational. Accounting liabilities are not investment
leverage, an annualized historical expense ratio is not a future fee promise,
and a trailing median spread is not a contemporaneous bid/ask quote.
"""
from copy import deepcopy, copy
from datetime import date
from math import isfinite, isclose, ceil
from pathlib import Path
import json
import re

POLICY = {'contract': 'commodity_trust_observed_exposure', 'version': '1',
    'accounting_basis': 'reported_net_asset_value',
    'investment_basis': 'physical_gold_value_before_accounting_liabilities',
    'liquidity_basis': 'separately_dated_volume_and_30_day_median_spread',
    'cost_basis': 'historical_period_annualized_observation', 'future_cost_basis': 'unavailable'}
CONTRACT = 'commodity_trust_observed_exposure/1'
PREPARATION_POLICY = (
    'Commodity trust observed exposure /1; no corporate valuation or future expense forecast. '
    'Exact opt-in policy: '+json.dumps(POLICY, sort_keys=True)+'. '
    'model only, scenarios={}, analysis_rationale required. All seven identity/holdings/terms/fees/risks/liquidity/quotation '
    'contracts require exact historical record_pointer proofs. Perimeter/calendar/policy are cited structural choices. '
    'holdings={assets:{gold},liabilities:{sponsor_fee_payable},reported:{total_assets,total_liabilities,net_assets,shares},as_of,currency}; '
    'assets minus all reported liabilities must reconcile to reported NAV without a plug. '
    'terms={replication:direct_asset,benchmark,effective_date,investment_leverage:null,leverage_basis:not_separately_established}. '
    'fees={period_start,period_end,sponsor_fees,total_expenses,annualized_expense_ratio,annualization_basis:'
    'reported_annualized_historical_expense_ratio,future_expense_ratio:null,other_future_fees_ratio:null,contingent_expenses}; '
    'record fees period is the exact reported duration. Historical annualized expense ratios are not future fee promises. '
    'liquidity={daily_volume,volume_unit:quoted units,spread_ratio,spread_basis:30_day_median_bid_ask,spread_window_days:30,observed_date,provider}; '
    'its period is observed_date, separately dated from accounting; median spread is not a spot spread. '
    'quotation retains its own price_as_of period. Identity and risks use the common observed exact shapes. '
    'Accounting gold/NAV and fee liability/NAV fractions are not investment leverage. '
    'Unknown future or contingent costs remain null, never zero; no geographic or derivative exposure is invented.'
)
SCHEMA = {
    'perimeter': ('instrument_identity', 'contract', 'scope', 'opening', 'contract', 'model'),
    'calendar': ('exposure_terms', 'contract', 'calendar', 'opening', 'contract', 'model'),
    'policy': ('exposure_terms', 'contract', 'observational_policy', 'opening', 'contract', 'model'),
    'identity': ('instrument_identity', 'contract', 'instrument', 'opening', 'contract', 'model'),
    'holdings': ('exposure_terms', 'contract', 'reported_nav_accounting', 'opening', 'contract', 'model'),
    'terms': ('exposure_terms', 'contract', 'physical_investment_scope', 'opening', 'contract', 'model'),
    'fees': ('fees_and_risks', 'contract', 'historical_period_cost', 'opening', 'contract', 'model'),
    'risks': ('fees_and_risks', 'contract', 'risk_disclosure', 'opening', 'contract', 'model'),
    'liquidity': ('fees_and_risks', 'contract', 'separately_dated_market_liquidity', 'opening', 'contract', 'model'),
    'quotation': ('quotation_units', 'contract', 'quotation', 'opening', 'contract', 'model'),
}


def selected(values):
    return isinstance(values, dict) and values.get('policy') == POLICY


def selected_records(records):
    policies = [row.get('value') for row in records if row.get('driver') == 'policy' and row.get('scenario') == 'model']
    return policies == [POLICY]


def record_period(driver, value, calendar):
    if driver == 'fees' and isinstance(value, dict):
        return str(value.get('period_start')) + '/' + str(value.get('period_end'))
    if driver == 'liquidity' and isinstance(value, dict):
        return value.get('observed_date')
    if driver == 'quotation' and isinstance(value, dict):
        return value.get('price_as_of')
    return calendar.get('valuation_date')


def _number(value):
    try:
        return type(value) in (int, float) and isfinite(value)
    except OverflowError:
        return False


def _date(value):
    try:
        result = date.fromisoformat(value)
        return result if result.isoformat() == value else None
    except (ValueError, TypeError):
        return None


def project(values, *, ticker, instrument, cutoff):
    issues = []
    def exact(name, keys):
        value = values.get(name)
        if not isinstance(value, dict) or set(value) != set(keys):
            issues.append(name + ': exact observed commodity-trust contract required')
            return {}
        return value
    perimeter = exact('perimeter', ('entity', 'currency', 'share_class'))
    calendar = exact('calendar', ('valuation_date', 'periods', 'discount_convention'))
    identity = exact('identity', ('ticker', 'name', 'instrument', 'issuer', 'legal_structure', 'underlying'))
    holdings = exact('holdings', ('assets', 'liabilities', 'reported', 'as_of', 'currency'))
    terms = exact('terms', ('replication', 'benchmark', 'effective_date', 'investment_leverage', 'leverage_basis'))
    fees = exact('fees', ('period_start', 'period_end', 'sponsor_fees', 'total_expenses', 'annualized_expense_ratio',
                         'annualization_basis', 'future_expense_ratio', 'other_future_fees_ratio', 'contingent_expenses'))
    risks = exact('risks', ('issuer', 'counterparty', 'custody', 'liquidity', 'tracking'))
    liquidity = exact('liquidity', ('daily_volume', 'volume_unit', 'spread_ratio', 'spread_basis', 'spread_window_days', 'observed_date', 'provider'))
    quote = exact('quotation', ('financial_currency', 'quote_currency', 'quote_unit', 'quote_units_per_currency',
                               'financial_to_quote_rate', 'shares_per_quote', 'share_class', 'price', 'price_as_of'))
    on, end = _date(calendar.get('valuation_date')), _date(cutoff)
    if not selected(values): issues.append('policy: exact versioned observed commodity-trust scope required')
    if not on or not end or on > end or calendar.get('periods') != [] or calendar.get('discount_convention') != 'snapshot':
        issues.append('calendar: a dated accounting snapshot before the cutoff is required')
    if perimeter and (any(not isinstance(v, str) or not v.strip() for v in perimeter.values())
                      or not re.fullmatch('[A-Z]{3}', perimeter.get('currency', ''))):
        issues.append('perimeter: exact legal entity, financial currency and entitled share class required')
    if identity and (identity['ticker'] != ticker or identity['instrument'] != instrument
                     or identity['name'] != perimeter.get('entity') or
                     any(not isinstance(v, str) or not v.strip() for v in identity.values())):
        issues.append('identity: observed name/ticker/instrument differs from the declared legal perimeter')
    assets, liabilities, reported = holdings.get('assets'), holdings.get('liabilities'), holdings.get('reported')
    if (not isinstance(assets, dict) or set(assets) != {'gold'} or not isinstance(liabilities, dict)
            or set(liabilities) != {'sponsor_fee_payable'} or not isinstance(reported, dict)
            or set(reported) != {'total_assets', 'total_liabilities', 'net_assets', 'shares'}):
        issues.append('holdings: complete observed asset/liability/NAV/share ledger required')
    elif (any(not _number(v) or v < 0 for v in [*assets.values(), *liabilities.values(), *reported.values()])
          or reported['net_assets'] <= 0 or reported['shares'] <= 0 or assets['gold'] <= 0):
        issues.append('holdings: finite nonnegative reported balances and positive NAV/shares/gold required')
    else:
        for key, calculated in [('total_assets', sum(assets.values())), ('total_liabilities', sum(liabilities.values())),
                                ('net_assets', sum(assets.values())-sum(liabilities.values()))]:
            if not isclose(calculated, reported[key], rel_tol=0, abs_tol=1e-7):
                issues.append('holdings: ' + key + ' does not reconcile; no balancing plug or omitted claim')
    if holdings and (holdings['as_of'] != calendar.get('valuation_date') or holdings['currency'] != perimeter.get('currency')):
        issues.append('holdings: balance date or currency differs from the accounting perimeter')
    if terms and (terms['replication'] != 'direct_asset' or terms['effective_date'] != calendar.get('valuation_date')
                  or terms['investment_leverage'] is not None or terms['leverage_basis'] != 'not_separately_established'
                  or not isinstance(terms['benchmark'], str) or not terms['benchmark'].strip()):
        issues.append('terms: physical scope must retain investment leverage as unavailable, not infer it from accounting weights')
    if fees:
        start, finish = _date(fees['period_start']), _date(fees['period_end'])
        if (not start or not finish or not on or not start <= finish == on or
                any(not _number(fees[k]) or fees[k] < 0 for k in ('sponsor_fees', 'total_expenses', 'annualized_expense_ratio'))
                or fees['total_expenses'] < fees['sponsor_fees'] or fees['annualized_expense_ratio'] > 1
                or fees['annualization_basis'] != 'reported_annualized_historical_expense_ratio'
                or fees['future_expense_ratio'] is not None or fees['other_future_fees_ratio'] is not None
                or not isinstance(fees['contingent_expenses'], str) or not fees['contingent_expenses'].strip()):
            issues.append('fees: exact historical duration and reported annualized ratio required; future and contingent costs remain unavailable')
    if risks and any(not isinstance(v, str) or not v.strip() for v in risks.values()):
        issues.append('risks: each primary disclosure must be explicit')
    if liquidity:
        observed = _date(liquidity['observed_date'])
        if (not observed or not end or observed > end or not _number(liquidity['daily_volume']) or liquidity['daily_volume'] < 0
                or liquidity['volume_unit'] != 'quoted units' or not _number(liquidity['spread_ratio'])
                or not 0 <= liquidity['spread_ratio'] <= 1 or liquidity['spread_basis'] != '30_day_median_bid_ask'
                or type(liquidity['spread_window_days']) is not int or liquidity['spread_window_days'] != 30
                or not isinstance(liquidity['provider'], str) or not liquidity['provider'].strip()):
            issues.append('liquidity: separately dated daily volume and trailing 30-day median spread required; not a spot spread')
    if quote:
        observed = _date(quote['price_as_of'])
        currency = perimeter.get('currency')
        if (not observed or not end or observed > end or any(not _number(quote[k]) or quote[k] <= 0 for k in
                ('price', 'quote_units_per_currency', 'financial_to_quote_rate', 'shares_per_quote'))
                or quote['financial_currency'] != currency or quote['quote_currency'] != currency or quote['quote_unit'] != currency
                or any(quote[k] != 1 for k in ('quote_units_per_currency', 'financial_to_quote_rate', 'shares_per_quote'))
                or quote['share_class'] != perimeter.get('share_class')):
            issues.append('quotation: exact separately dated same-currency ordinary trust share observation required')
    if issues: return {'status': 'incomplete', 'issues': issues}
    nav, gold, liability = reported['net_assets'], assets['gold'], liabilities['sponsor_fee_payable']
    gold_ratio, liability_ratio = gold/nav, -liability/nav
    try:
        squared = gold_ratio**2+liability_ratio**2
        nav_per_share = nav/reported['shares']
        gross, net = abs(gold_ratio)+abs(liability_ratio), gold_ratio+liability_ratio
        if not all(_number(value) for value in (gold_ratio, liability_ratio, squared, nav_per_share, gross, net)):
            raise OverflowError('Nonfinite derived metric')
    except OverflowError:
        return {'status': 'incomplete', 'issues': ['Accounting NAV ratios and share/concentration results must be finite and representable']}
    positions = [{'id': 'reported_gold', 'name': 'Reported physical gold', 'weight': gold_ratio,
                  'asset_class': 'physical_gold', 'country': None, 'currency': perimeter['currency']},
                 {'id': 'reported_fee_payable', 'name': 'Reported sponsor fee payable', 'weight': liability_ratio,
                  'asset_class': 'accounting_liability', 'country': None, 'currency': perimeter['currency']}]
    return {'status': 'complete', 'issues': [], 'contract': CONTRACT, 'identity': deepcopy(identity),
        'accounting': deepcopy(holdings), 'accounting_date': calendar['valuation_date'],
        'terms': deepcopy(terms), 'fees': deepcopy(fees), 'risks': deepcopy(risks), 'liquidity': deepcopy(liquidity),
        'quotation': deepcopy(quote), 'positions': positions, 'investment_leverage': None,
        'gold_nav_ratio': gold_ratio, 'liability_nav_ratio': liability_ratio, 'net_exposure': net,
        'gross_exposure': gross, 'largest_absolute_weight': abs(gold_ratio),
        'squared_weight_sum': squared, 'annual_declared_cost_ratio': None,
        'historical_annualized_expense_ratio': fees['annualized_expense_ratio'], 'accounting_nav_per_share': nav_per_share,
        'buckets': {'asset_class': [{'label': p['asset_class'], 'weight': p['weight']} for p in positions],
                    'currency': [{'label': perimeter['currency'], 'weight': gold_ratio+liability_ratio}], 'country': []},
        'concentration_basis': 'Accounting fractions of reported NAV. Gold and accrued fee liability are distinct; absolute balance-sheet weights are not investment leverage. Geographic classification is unavailable.',
        'limitations': ['Investment leverage is not separately established.', 'Future and contingent expense ratios are unavailable.',
                       'The trailing 30-day median spread is not a spot bid/ask quote.',
                       'Accounting, issuer price and market liquidity retain their separate observed dates.']}


def bind(bundle):
    from .dcf_quality import _evidence_issues
    case = bundle['case']; records = case['records']; values = {}; issues = []; consumed = []; evidence = []
    for index, record in enumerate(records):
        driver = record.get('driver'); errors = []
        if driver not in SCHEMA or driver in values or record.get('scenario') != 'model':
            issues.append({'field': str(driver), 'reason': 'Duplicate, unconsumed or non-model commodity observation'}); continue
        field, unit, basis, *_ = SCHEMA[driver]
        calendar = next((row['value'] for row in records if row.get('driver') == 'calendar'), {})
        perimeter = next((row['value'] for row in records if row.get('driver') == 'perimeter'), {})
        expected = {'field': field, 'unit': unit, 'accounting_basis': basis, 'entity': perimeter.get('entity'),
                    'period': record_period(driver, record.get('value'), calendar)}
        errors.extend(key + ' differs from the observed contract' for key, value in expected.items() if record.get(key) != value)
        historical = driver not in ('perimeter', 'calendar', 'policy')
        if record.get('kind') != ('historical' if historical else 'analyst_estimate'):
            errors.append('Historical observation or explicit structural policy kind required')
        proof = {'kind': record.get('kind'), 'source': record.get('source_id'), 'source_date': record.get('as_of'),
                 'valid_until': record.get('valid_until'), 'rationale': record.get('rationale'), 'metric': driver, 'basis': basis}
        errors.extend(_evidence_issues(proof, _date(case['as_of'])))
        if errors:
            issues.append({'field': field, 'reason': driver + ': ' + '; '.join(errors)}); continue
        values[driver] = deepcopy(record['value'])
        consumed.append({'record_index': index, **{key: record[key] for key in ('field','scenario','driver','entity','period','source_id')}})
        evidence.append({'scenario': 'model', 'driver': driver, 'values': deepcopy(record['value']), 'evidence': proof, 'issues': []})
    if set(values) != set(SCHEMA): issues.append({'field': 'exposure_terms', 'reason': 'All ten exact commodity-trust contracts must be consumed'})
    if case['assumptions'] or set(bundle['analysis_context']) != {'analysis_rationale'}:
        issues.append({'field': 'analysis_context', 'reason': 'No legacy assumptions or unconsumed analysis context allowed'})
    return {'values': {'model': values}, 'issues': issues, 'consumed': consumed, 'evidence': evidence,
            'calendar': values.get('calendar') or {}, 'perimeter': values.get('perimeter') or {}, 'quotation': values.get('quotation') or {}}


def assemble_sources(bundle, report, *, as_of):
    """A server-compiled primary catalog may supply observed contracts, no plan."""
    from .input_preparation import _catalog, _day, prepare_method_inputs
    from .preparation_fresh_historical import _issuer_key
    from .preparation_commodity_sources import assemble_commodity_catalog
    result=assemble_commodity_catalog(bundle,report,as_of=as_of)
    receipt={'contract':'commodity_trust_historical_sources/1','status':'blocked','reasons':[],
        'historical_drivers':[],'forecast_values_reused':False,'archive_reads':'none_by_contract',
        'limitation':'Acquired raw primary catalog is normalized automatically; new source discovery is not performed. Public projection requires its separate proof gate.'}
    result['commodity_trust_historical_assembly']=receipt
    candidates=[doc for doc in result.get('documents') or [] if isinstance(doc,dict) and
                (doc.get('metadata') or {}).get('normalizer')=='commodity_trust_primary/1']
    if not candidates:
        receipt['reasons']=result.get('commodity_trust_raw_catalog',{}).get('reasons') or ['No server-compiled commodity-trust primary observations in this catalog']
        return result
    try:
        catalog,issues,_=_catalog(result['documents'],_day(as_of))
        if issues:raise ValueError('; '.join(row['reason'] for row in issues))
        observed=[doc for doc in catalog.values() if (doc.get('metadata') or {}).get('normalizer')=='commodity_trust_primary/1']
        if len(observed)!=1:raise ValueError('Unique recompiled commodity-trust observation inventory required')
        document=observed[0];metadata=document['metadata'];entity=metadata['entity'];on=metadata['on']
        profile=bundle['case'].get('info') or {};confirmed=profile.get('longName') or profile.get('shortName')
        if metadata['ticker']!=bundle['case']['ticker'] or _issuer_key(entity)!=_issuer_key(confirmed):
            receipt['fatal_identity_mismatch']=True
            raise ValueError('Commodity-trust primary legal entity differs from the confirmed security profile')
        observations=json.loads(document['text'])['observations']
        required=set(SCHEMA)-{'perimeter','calendar','policy'}
        if len(observations)!=len(required) or {row.get('driver') for row in observations}!=required:
            raise ValueError('All seven unique primary observed commodity-trust contracts required')
        plan={'model':{},'scenarios':{},'analysis_rationale':'Reported commodity-trust accounting snapshot and separately dated public market observations; no corporate fair value or future expense forecast.'}
        for index,observation in enumerate(observations):
            driver=observation['driver']
            plan['model'][driver]={'value':deepcopy(observation['value']),'kind':'historical','evidence_ids':[document['id']],
                'record_pointer':'/observations/'+str(index),'rationale':'Exact primary commodity-trust observation, independently recompiled from original archived bytes.',
                'valid_until':as_of,'valid_until_basis':{'policy':'same_day','as_of':as_of}}
        quote=plan['model']['quotation']['value']
        structures={'perimeter':{'entity':entity,'currency':quote['financial_currency'],'share_class':quote['share_class']},
                    'calendar':{'valuation_date':on,'periods':[],'discount_convention':'snapshot'},'policy':deepcopy(POLICY)}
        for driver,value in structures.items():
            plan['model'][driver]={'value':value,'kind':'analyst_estimate','evidence_ids':[metadata['source_document_id']],
                'rationale':'Explicit dated observed commodity-trust scope; no rolling accounting balances, investment leverage or prospective fee assumptions.',
                'valid_until':as_of,'valid_until_basis':{'policy':'same_day','as_of':as_of}}
        projection=project({driver:item['value'] for driver,item in plan['model'].items()},
            ticker=bundle['case']['ticker'],instrument=bundle['decision']['instrument'],cutoff=as_of)
        if projection['status']!='complete':raise ValueError('; '.join(projection['issues']))
        checked=prepare_method_inputs(bundle,documents=list(catalog.values()),source_report={**result,'source_plan':plan},propose=lambda *_:deepcopy(plan))
        if checked['status']!='prepared':raise ValueError('; '.join(row['reason'] for row in checked['issues']))
        receipt.update(status='ready',historical_drivers=sorted(required),economic_opening_date=on,information_cutoff=as_of,
            separately_dated_observations={driver:record_period(driver,item['value'],structures['calendar']) for driver,item in plan['model'].items()},
            limitations=projection['limitations'])
        used=set(metadata['source_document_ids'].values())|{document['id']}
        result.update(source_plan=plan,documents=list(catalog.values()),preparation_ready=True,fresh_historical_assembly=deepcopy(receipt),
            selection={'policy':'verified_commodity_trust_observations','opening_date':on,
                       'selected_document_ids':sorted(used),'excluded_document_ids':sorted(set(catalog)-used)})
    except (ValueError,TypeError,KeyError,IndexError,OSError) as exc:
        receipt['reasons'].append(str(exc))
        if receipt.get('fatal_identity_mismatch'):result['fresh_historical_assembly']=deepcopy(receipt)
    return result


def leaves(value, path=()):
    if isinstance(value, dict):
        for key in sorted(value): yield from leaves(value[key], path+(key,))
    elif isinstance(value, list):
        for index, child in enumerate(value): yield from leaves(child, path+(index,))
    else: yield path, value


def _excel_value(value):
    # Match the common XLSX writer's numeric representation exactly; this is
    # serialization precision, not a tolerance for changed observations.
    if value is None: return 'n.d.'
    return float(format(value, '.16g')) if type(value) is float else value


def _observed_format(value, *, ratio=False, count=False):
    if _number(value) and 0 < abs(value) < .000001:
        return '0.000000E+00'
    if ratio or _number(value) and 0 < abs(value) < 1:
        return '0.000000'
    if count and _number(value):
        return '#,##0' if value == int(value) else '#,##0.000000'
    return '#,##0.00'


def _observed_layout(sheet, widths):
    """Retain full dollar values and source disclosures on screen and in print."""
    for column,width in widths.items(): sheet.column_dimensions[column].width=width
    for row in sheet.iter_rows(min_row=8):
        lines=1
        for cell in row:
            if not isinstance(cell.value,str) or cell.data_type=='f': continue
            alignment=copy(cell.alignment);alignment.wrap_text=True;alignment.vertical='top';cell.alignment=alignment
            width=sheet.column_dimensions[cell.column_letter].width
            lines=max(lines,sum(max(1,ceil(len(part)/max(1,width*.9-2))) for part in cell.value.split('\n')))
        sheet.row_dimensions[row[0].row].height=min(409,max(21,lines*15+6))


def workbook(payload, output_dir):
    from openpyxl import Workbook
    from openpyxl.workbook.properties import CalcProperties
    from .documented_presentation import _sheet, _line, _header, _put, _finish, PRICE
    from .documented_formulas import _formula
    from .dcf_quality_sheet import append_quality_sheet, append_sector_quality_sheet
    from .sourcebook_presentation import present_sourcebook
    directory = Path(output_dir); directory.mkdir(parents=True, exist_ok=True)
    path = directory / ('EXPOSURE_' + re.sub(r'[^A-Za-z0-9_-]', '_', payload['ticker']) + '_' + payload['generation_id'] + '.xlsx')
    wb = Workbook(); wb.remove(wb.active)
    inputs = _sheet(wb, 'Observed inputs', 'Dated public commodity-trust observations', 7)
    _header(inputs, 7, ['Driver','Path','Observed value','Unit','Period'])
    refs, cells = [], {}; row = 8
    for record in sorted(payload['acquisition_snapshot']['case']['records'], key=lambda item: item['driver']):
        for leaf, value in leaves(record['value']):
            _put(inputs,row,2,record['driver']); _put(inputs,row,3,'.'.join(map(str,leaf)))
            _put(inputs,row,4,_excel_value(value)); _put(inputs,row,5,record['unit']); _put(inputs,row,6,record['period'])
            if _number(value):
                last = leaf[-1] if leaf else ''
                inputs.cell(row,4).number_format = _observed_format(value,
                    ratio=last in ('annualized_expense_ratio','spread_ratio'), count=last in ('daily_volume','shares','spread_window_days'))
            cell = "'Observed inputs'!D"+str(row); cells[(record['driver'],leaf)] = cell
            if _number(value): refs.append({'scenario':'model','driver':record['driver'],'path':list(leaf),'cell':cell})
            row += 1
    _finish(inputs,row+1,7)
    _observed_layout(inputs,{'B':24,'C':42,'D':56,'E':16,'F':30})
    p = payload['exposure_analysis']; get = lambda driver,*path: cells[(driver,path)]
    nav = get('holdings','reported','net_assets'); gold = get('holdings','assets','gold')
    liability = get('holdings','liabilities','sponsor_fee_payable')
    expressions = [
        ('calculated_net_assets','Calculated reported NAV','USD',gold+'-'+liability,p['accounting']['reported']['net_assets']),
        ('reported_net_assets','Reported NAV','USD',nav,p['accounting']['reported']['net_assets']),
        ('nav_reconciliation','NAV residual','USD','D8-D9',0.),
        ('accounting_nav_per_share','Accounting NAV per share','USD per share','D8/'+get('holdings','reported','shares'),p['accounting_nav_per_share']),
        ('gold_nav_ratio','Physical gold / reported NAV','ratio',gold+'/D8',p['gold_nav_ratio']),
        ('liability_nav_ratio','Accrued sponsor fee / NAV','ratio','-'+liability+'/D8',p['liability_nav_ratio']),
        ('net_exposure','Net accounting weights','ratio','SUM(D12:D13)',p['net_exposure']),
        ('asset_reconciliation','Asset residual','USD',gold+'-'+get('holdings','reported','total_assets'),0.),
        ('liability_reconciliation','Liability residual','USD',liability+'-'+get('holdings','reported','total_liabilities'),0.),
        ('historical_annualized_expense_ratio','Historical annualized expense ratio','ratio',get('fees','annualized_expense_ratio'),p['historical_annualized_expense_ratio']),
        ('median_spread_ratio','Trailing 30-day median spread','ratio',get('liquidity','spread_ratio'),p['liquidity']['spread_ratio']),
        ('daily_volume','Separately dated daily volume','quoted units',get('liquidity','daily_volume'),p['liquidity']['daily_volume']),
        ('price','Separately dated issuer price',payload['currency'],get('quotation','price'),payload['price'])]
    summary = _sheet(wb,'Summary',payload['ticker']+' | Observed commodity trust',6)
    _line(summary,4,'Historical NAV accounting; market observations retain separate dates. No corporate fair value, future fee promise or inferred investment leverage.',6,height=44)
    _header(summary,7,['Measure','Unit','Value','Economic period'])
    baseline=wb.create_sheet('Engine Baseline');baseline.append(['Engine path','Recorded value'])
    metrics={}; formulas={}
    for row,(key,label,unit,formula,value) in enumerate(expressions,8):
        _put(summary,row,2,label);_put(summary,row,3,unit)
        _formula(summary,row,4,formula,_observed_format(value,ratio=unit=='ratio',count=key=='daily_volume'))
        period = record_period('fees',p['fees'],{}) if key=='historical_annualized_expense_ratio' else p['liquidity']['observed_date'] if key in ('median_spread_ratio','daily_volume') else p['quotation']['price_as_of'] if key=='price' else p['accounting_date']
        _put(summary,row,5,period);baseline.append([key,value]);metrics[key]="'Summary'!D"+str(row);formulas[key]=formula
    for row,label,value in ((23,'Investment leverage','n.d. — not separately established'),(24,'Future annual fees','n.d. — future and contingent costs unavailable'),
                            (25,'Accounting date',p['accounting_date']),(26,'Liquidity observation',p['liquidity']['observed_date']),(27,'Issuer price date',p['quotation']['price_as_of'])):
        _put(summary,row,2,label);_put(summary,row,4,value)
    _finish(summary,29,6)
    _observed_layout(summary,{'B':50,'C':20,'D':28,'E':30})
    checks=_sheet(wb,'Observation checks','Historical accounting checks',5);_header(checks,7,['Control','Result'])
    for row,key in enumerate(('nav_reconciliation','asset_reconciliation','liability_reconciliation'),8):
        _put(checks,row,2,key);_formula(checks,row,3,'IF(ABS('+metrics[key]+')<=0.0000001,"OK","KO")')
    _finish(checks,12,5)
    append_quality_sheet(wb,payload['analytical_quality']);append_sector_quality_sheet(wb,payload);present_sourcebook(wb,payload)
    wb.move_sheet(summary,-wb.index(summary));wb.active=0;wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True)
    payload['calculation_details']['workbook_bindings']={'contract':'commodity_trust_workbook/1','input_refs':refs,'metrics':metrics,'formulas':formulas,
        'price':metrics['price'],'observation_cells':{'accounting_date':"'Summary'!D25",'liquidity_date':"'Summary'!D26",'price_date':"'Summary'!D27"}}
    wb.save(path);wb.close();return str(path)


def exhibits(payload, workbook, result):
    p=payload['exposure_analysis'];packet=payload['calculation_details']['workbook_bindings']
    if (payload.get('price')!=p['quotation']['price'] or payload.get('currency')!=p['quotation']['quote_unit']
            or payload.get('financial_currency')!=p['accounting']['currency'] or payload.get('valuation_date')!=p['accounting_date']):
        raise ValueError('Commodity headline observation identity/date/currency differs from acquired contracts')
    if packet.get('contract')!='commodity_trust_workbook/1': raise ValueError('Exact commodity observation workbook contract required')
    expected_refs=[];cells={};row=8
    for record in sorted(payload['acquisition_snapshot']['case']['records'],key=lambda item:item['driver']):
        for leaf,value in leaves(record['value']):
            sheet=workbook['Observed inputs']
            expected=[record['driver'],'.'.join(map(str,leaf)),_excel_value(value),record['unit'],record['period']]
            if [sheet.cell(row,col).value for col in range(2,7)]!=expected:
                raise ValueError('Commodity workbook observed value/period mismatch: '+record['driver'])
            if _number(value):
                last=leaf[-1] if leaf else ''
                if sheet.cell(row,4).number_format != _observed_format(value,
                        ratio=last in ('annualized_expense_ratio','spread_ratio'), count=last in ('daily_volume','shares','spread_window_days')):
                    raise ValueError('Commodity observed value display precision differs: '+record['driver'])
                expected_refs.append({'scenario':'model','driver':record['driver'],'path':list(leaf),'cell':"'Observed inputs'!D"+str(row)})
            cells[(record['driver'],leaf)]="'Observed inputs'!D"+str(row)
            row+=1
    if any(workbook['Observed inputs'].cell(row,col).value is not None for col in range(2,7)):
        raise ValueError('Unbound extra commodity observation')
    if expected_refs!=packet['input_refs']:raise ValueError('Commodity workbook bindings differ from acquired observations')
    baseline={row[0].value:row[1].value for row in workbook['Engine Baseline'] if len(row)>=2}
    expected_values={'calculated_net_assets':p['accounting']['reported']['net_assets'],'reported_net_assets':p['accounting']['reported']['net_assets'],
        'nav_reconciliation':0.,'accounting_nav_per_share':p['accounting_nav_per_share'],'gold_nav_ratio':p['gold_nav_ratio'],'liability_nav_ratio':p['liability_nav_ratio'],
        'net_exposure':p['net_exposure'],'asset_reconciliation':0.,'liability_reconciliation':0.,'historical_annualized_expense_ratio':p['historical_annualized_expense_ratio'],
        'median_spread_ratio':p['liquidity']['spread_ratio'],'daily_volume':p['liquidity']['daily_volume'],'price':payload['price']}
    get=lambda driver,*path:cells[(driver,path)]
    gold=get('holdings','assets','gold');liability=get('holdings','liabilities','sponsor_fee_payable');nav=get('holdings','reported','net_assets')
    expected_formulas=dict(zip(expected_values,(gold+'-'+liability,nav,'D8-D9','D8/'+get('holdings','reported','shares'),gold+'/D8','-'+liability+'/D8',
        'SUM(D12:D13)',gold+'-'+get('holdings','reported','total_assets'),liability+'-'+get('holdings','reported','total_liabilities'),
        get('fees','annualized_expense_ratio'),get('liquidity','spread_ratio'),get('liquidity','daily_volume'),get('quotation','price'))))
    expected_metrics={key:"'Summary'!D"+str(index) for index,key in enumerate(expected_values,8)}
    if packet['metrics']!=expected_metrics or packet['formulas']!=expected_formulas or packet.get('price')!=expected_metrics['price']:
        raise ValueError('Commodity decisive formulas/bindings differ from the common layout')
    for key,value in expected_values.items():
        row=int(packet['metrics'][key].split('D')[-1])
        period=record_period('fees',p['fees'],{}) if key=='historical_annualized_expense_ratio' else p['liquidity']['observed_date'] if key in ('median_spread_ratio','daily_volume') else p['quotation']['price_as_of'] if key=='price' else p['accounting_date']
        if (baseline.get(key)!=_excel_value(value) or workbook['Summary'].cell(row,4).value!='='+expected_formulas[key]
                or workbook['Summary'].cell(row,5).value!=period):
            raise ValueError('Commodity engine/workbook formula or baseline mismatch: '+key)
        ratio=key in ('gold_nav_ratio','liability_nav_ratio','net_exposure','historical_annualized_expense_ratio','median_spread_ratio')
        if workbook['Summary'].cell(row,4).number_format != _observed_format(value,ratio=ratio,count=key=='daily_volume'):
            raise ValueError('Commodity decisive observation display precision differs: '+key)
    for row,value in ((23,'n.d. — not separately established'),(24,'n.d. — future and contingent costs unavailable'),
                       (25,p['accounting_date']),(26,p['liquidity']['observed_date']),(27,p['quotation']['price_as_of'])):
        if workbook['Summary'].cell(row,4).value!=value:raise ValueError('Commodity unavailable value/date claim differs')
    for row,key in enumerate(('nav_reconciliation','asset_reconciliation','liability_reconciliation'),8):
        if workbook['Observation checks'].cell(row,3).value!='=IF(ABS('+expected_metrics[key]+')<=0.0000001,"OK","KO")':
            raise ValueError('Commodity reconciliation control differs')
    accounting_units={'calculated_net_assets':payload['financial_currency'],'reported_net_assets':payload['financial_currency'],
        'nav_reconciliation':payload['financial_currency'],'accounting_nav_per_share':payload['financial_currency']+' per entitled share',
        'gold_nav_ratio':'fraction of reported NAV','liability_nav_ratio':'fraction of reported NAV','net_exposure':'fraction of reported NAV'}
    market_units={'historical_annualized_expense_ratio':'annualized historical fraction','median_spread_ratio':'30-day median fraction',
                  'daily_volume':'quoted units','price':payload['currency']+' per quoted unit'}
    rows=[[key,expected_values[key],accounting_units[key]] for key in accounting_units]
    market=['historical_annualized_expense_ratio','median_spread_ratio','daily_volume','price']
    result.update(status='complete',exhibits=[
        {'id':'commodity_accounting','title':'Reported NAV accounting','kind':'table','columns':['Measure','Value','Unit'],
         'rows':rows,'cell_refs':[[None,packet['metrics'][r[0]],None] for r in rows],'unit':'mixed','currency':None,
         'period':p['accounting_date'],'scenario':'observed','source':'exact_common_engine_and_workbook','explanation':p['concentration_basis']},
        {'id':'commodity_market_costs','title':'Dated market and historical cost observations','kind':'table','columns':['Measure','Value','Unit','Period'],
         'rows':[[key,expected_values[key],market_units[key],record_period('fees',p['fees'],{}) if key.startswith('historical') else p['quotation']['price_as_of'] if key=='price' else p['liquidity']['observed_date']] for key in market],
         'cell_refs':[[None,packet['metrics'][key],None,None] for key in market],'unit':'mixed','currency':None,'period':'separately_dated',
         'scenario':'observed','source':'exact_common_engine_and_workbook','explanation':'; '.join(p['limitations'])}],
        input_bindings=deepcopy(packet['input_refs']),intrinsic_value_applicable=False,analysis_usability=deepcopy(payload['analysis_usability']),
        unavailable_exhibits=[{'id':key,'reason':'Corporate intrinsic value is not applicable to an observed commodity trust'} for key in ('valuation_bridge','scenarios','sensitivity','implicit_expectations')]
            +[{'id':'investment_leverage','reason':p['limitations'][0]},{'id':'future_expenses','reason':p['limitations'][1]}])
    result['headline_values']={'price':{'value':payload['price'],'cell':packet['price'],'unit':'per quoted unit','currency':payload['currency'],'period':p['quotation']['price_as_of']}}
    return result
