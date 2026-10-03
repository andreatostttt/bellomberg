"""Read-only workbook view of already acquired financial, consensus and guidance data.

No provider is called here. Raw units and provider period labels are deliberately
preserved: a displayed history is not a certified comparable time series.
"""
from math import ceil, isfinite
import json
from urllib.parse import urlsplit
from openpyxl.styles import Alignment
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.utils import get_column_letter

from bellomberg.core.language import text as tr

from .documented_presentation import _finish, _header, _line, _put, _sheet

# General retains small ratios and lets Excel use scientific notation for raw magnitudes.
AMOUNT = 'General'


MISSING = 'n.d.'
SOURCE_NAMES = ('financials', 'consensus', 'guidance')
_METADATA = frozenset(('source_id', 'as_of', 'status', 'unit', 'valid_until'))


def _source_text(value):
    if value is None or value == '':
        return MISSING
    if isinstance(value, str):
        return value
    try:
        return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False)
    except (ValueError, TypeError):
        return tr('n.d. (valore non rappresentabile)', 'n.a. (value not representable)')


def _kind(kind):
    labels = {'historical': ('Fatto riportato', 'Reported fact'),
              'company_guidance': ('Guidance societaria', 'Company guidance'),
              'analyst_estimate': ('Stima analista', 'Analyst estimate'),
              'proxy': ('Proxy dichiarato', 'Declared proxy')}
    return tr(*labels[kind]) if kind in labels else tr('Tipo n.d.', 'Type n.a.')


def _consumed_source_rows(payload):
    """Join exact consumed records to their compiled plan; never borrow old proof."""
    case = (payload.get('acquisition_snapshot') or {}).get('case') or {}
    records = case.get('records') or []
    preparation = payload.get('preparation') or {}
    proposal = preparation.get('proposal') or {}
    plan = proposal.get('plan') or {}
    seed = (preparation.get('review_basis') or {}).get('seed') or {}
    plan_bound = False
    if seed.get('plan_sha256'):
        from .preparation_seed import _digest
        try:
            plan_bound = seed['plan_sha256'] == _digest(plan) == _digest(seed.get('plan'))
        except (ValueError, TypeError):
            pass
    documents = (preparation.get('provenance') or {}).get('documents') or {}
    originals = {doc.get('id'): doc for doc in
                 ((preparation.get('review_basis') or {}).get('dossier') or {}).get('documents') or []
                 if isinstance(doc, dict)}
    identity = ('field', 'scenario', 'driver', 'entity', 'period', 'source_id')
    contract = (*identity, 'value', 'unit', 'accounting_basis', 'source_locator',
                'as_of', 'valid_until', 'kind', 'rationale')
    for receipt in (payload.get('input_consumption') or {}).get('consumed_records') or []:
        index = receipt.get('record_index')
        record = records[index] if type(index) is int and 0 <= index < len(records) else None
        matched = isinstance(record, dict) and all(record.get(key) == receipt.get(key) for key in identity)
        if not matched:
            yield {'key': (receipt.get('scenario'), receipt.get('driver')), 'kind': None,
                'values': [receipt.get('scenario'), receipt.get('driver'), _kind(None), MISSING,
                    MISSING, MISSING, MISSING, MISSING, MISSING, MISSING,
                    tr('Ricevuta di consumo non corrispondente', 'Consumption receipt does not match')]}
            continue
        scope, driver = record['scenario'], record['driver']
        candidates = [row for row in proposal.get('method_records') or [] if isinstance(row, dict)
                      and _source_text({key: row.get(key) for key in contract})
                      == _source_text({key: record.get(key) for key in contract})]
        item = ((plan.get('model') or {}).get(driver) if scope == 'model' else
                ((plan.get('scenarios') or {}).get(scope) or {}).get(driver))
        ids = [part.strip() for part in (record.get('source_locator') or '').split(',') if part.strip()]
        linked = plan_bound and len(candidates) == 1 and isinstance(item, dict) and item.get('evidence_ids') == ids
        proof = None
        if linked:
            fields = ('evidence_pointer', 'evidence_quote', 'quoted_value', 'quoted_unit',
                      'period_quote', 'record_pointer', 'calculation', 'facts', 'dilution_estimate')
            proof = {key: item[key] for key in fields if key in item}
            rationale = record.get('rationale') or ''
            if rationale.startswith('SOURCE POINTER NORMALIZED: '):
                try:
                    proof['compiler_pointer_normalizations'] = json.loads(rationale.split('\n', 1)[0].split(': ', 1)[1])
                except (ValueError, IndexError):
                    linked = False
                    proof = None
        binding = (tr('Record consumato; prova del piano collegata', 'Consumed record; compiled plan proof linked')
                   if linked else tr('Record consumato; Piano non collegato al record consumato',
                                     'Consumed record; plan is not linked to this consumed record'))
        for ident in ids or [None]:
            document = documents.get(ident) if ident is not None else None
            # A document ID identifies provenance; a raw provider envelope never substitutes it.
            original = originals.get(ident)
            inconsistent = isinstance(document, dict) and (
                bool(ids) and ident == ids[0] and document.get('url') != record.get('source_id')
                or isinstance(original, dict) and any(document.get(key) != original.get(key)
                    for key in ('url', 'published_at', 'sha256', 'document_sha256')
                    if document.get(key) is not None and original.get(key) is not None))
            if inconsistent:
                document = None
            url = (document.get('url') if isinstance(document, dict) else
                   record.get('source_id') if not ids or ident == ids[0] else None)
            missing = not isinstance(document, dict) or not document.get('url')
            status = binding + ('; '+tr('Metadati documento n.d.', 'Document metadata n.a.') if missing else '')
            if inconsistent:
                status += '; '+tr('Metadati documento incoerenti; URL dichiarato nel record',
                                   'Document metadata inconsistent; URL declared in the record')
            elif missing and url:
                status += '; '+tr('URL dichiarato nel record', 'URL declared in the record')
            dates = (tr('Pubblicazione: ', 'Publication: ')+_source_text(document.get('published_at') if isinstance(document, dict) else None)
                + '\n'+tr('Record osservato: ', 'Record observed: ')+_source_text(record.get('as_of'))
                + '\n'+tr('Valido fino: ', 'Valid until: ')+_source_text(record.get('valid_until')))
            pointer = _source_text(proof) if proof else tr('Pointer/prova n.d.; vedi motivazione e riferimenti dichiarati.',
                                                         'Pointer/proof n.a.; see rationale and declared references.')
            yield {'key': (scope, driver), 'kind': record.get('kind'),
                'values': [scope, driver, _kind(record.get('kind')), _source_text(record.get('value')),
                    ' | '.join(_source_text(record.get(key)) for key in ('unit', 'period', 'entity')),
                    _source_text(record.get('rationale')), ident or MISSING, dates, _source_text(url), pointer, status]}


def _text_parts(value):
    """Respect Excel's UTF-16 cell limit without dropping any source text."""
    text = _source_text(value)
    parts, current, units = [], [], 0
    for character in text:
        width = 2 if ord(character) > 0xffff else 1
        if units+width > 30000:
            parts.append(''.join(current)); current, units = [], 0
        current.append(character); units += width
    return [*parts, ''.join(current)]


def _http_link(cell):
    try:
        url = urlsplit(cell.value)
        if url.scheme in ('http', 'https') and url.netloc:
            cell.hyperlink = cell.value
    except (TypeError, ValueError):
        pass


def _assumption_links(wb, entries, anchors):
    mapping = getattr(wb, '_assumption_source_rows', None)
    if not mapping or 'Assumptions' not in wb:
        return
    ws = wb['Assumptions']; end = mapping['end']
    kinds = {entry['key']: _kind(entry['kind']) for entry in entries}
    for row in (7, 19, 31, 43):
        _header(ws, row, [tr('Tipo di dato', 'Data type'), tr('Motivazione e fonti', 'Rationale and sources')], start=end+1)
    for row, keys in mapping['rows'].items():
        labels = [key[0]+': '+kinds.get(key, tr('Tipo n.d.', 'Type n.a.')) for key in keys]
        _literal(ws, row, end+1, '\n'.join(labels))
        targets = [anchors[key] for key in keys if key in anchors]
        link = _literal(ws, row, end+2, tr('Apri motivazione e fonti consumate', 'Open rationale and consumed sources')
                        if targets else tr('Fonti consumate n.d.', 'Consumed sources n.a.'))
        if targets:
            for cell in (link, ws.cell(row, 2)):
                cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'Sources'!B{targets[0]}")
        for column in (end+1, end+2):
            ws.cell(row, column).alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 21, 18*len(keys)+6)
    ws.column_dimensions[get_column_letter(end+1)].width = 28
    ws.column_dimensions[get_column_letter(end+2)].width = 36
    ws.cell(mapping['footer'], 2).hyperlink = Hyperlink(ref=f"B{mapping['footer']}", location="'Sources'!B7")
    _finish(ws, mapping['footer']+1, end+2)


def _present_consumed_sources(wb, payload):
    entries = list(_consumed_source_rows(payload))
    if not entries:
        return False
    ws = _sheet(wb, 'Sources', tr('Fonti e motivazioni degli input consumati', 'Sources and rationales of consumed inputs'), 12)
    _line(ws, 4, tr('Snapshot documentale in sola lettura: fatti riportati, guidance e stime restano distinti. Nessuna approvazione economica o verifica live implicita.',
                    'Read-only documentary snapshot: reported facts, guidance and estimates remain distinct. No economic approval or live verification is implied.'), 12, height=40)
    _header(ws, 7, ['Scenario', 'Driver', tr('Tipo', 'Type'), tr('Valore compilato', 'Compiled value'),
        tr('Unita / periodo / entita', 'Unit / period / entity'), tr('Motivazione integrale', 'Full rationale'),
        tr('ID / riferimento documento', 'Document ID / reference'), tr('Date', 'Dates'), 'URL',
        tr('Pointer / prova dichiarata', 'Pointer / declared proof'), tr('Stato del legame', 'Binding status')])
    anchors, row = {}, 8
    for entry in entries:
        anchors.setdefault(entry['key'], row)
        parts = [_text_parts(value) for value in entry['values']]
        for part in range(max(map(len, parts))):
            for column, chunks in enumerate(parts, 2):
                value = chunks[part] if part < len(chunks) else (chunks[0] if column in (2, 3, 4, 8) else '')
                cell = _literal(ws, row, column, value)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if column == 10:
                    _http_link(cell)
            ws.row_dimensions[row].height = 90 if part == 0 else 120
            row += 1
    for column, width in {'B': 12, 'C': 28, 'D': 22, 'E': 36, 'F': 38, 'G': 70,
                          'H': 42, 'I': 32, 'J': 62, 'K': 75, 'L': 56}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = 'E8'
    ws.auto_filter.ref = f'B7:L{row-1}'
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    _finish(ws, row-1, 12)
    _assumption_links(wb, entries, anchors)
    if 'Assumptions' in wb:
        wb.move_sheet(ws, wb.index(wb['Assumptions'])+1-wb.index(ws))
    return True


def _literal(ws, row, col, value, fmt=None):
    """Provider text is never a formula, including leading =, +, - or @."""
    cell = _put(ws, row, col, value, fmt)
    if isinstance(value, str):
        cell.data_type = 's'
    return cell


def _unit(source, value=None):
    explicit = (value.get('unit') if isinstance(value, dict) else None) or source.get('unit')
    return explicit if isinstance(explicit, str) and explicit.strip() else MISSING


def _metadata(ws, source, as_of):
    for row, (name, value) in enumerate((
        (tr('Fonte', 'Source'), source.get('source_id') or MISSING),
        (tr('Osservata il', 'Observed on'), source.get('as_of') or MISSING),
        (tr('Cutoff', 'Cutoff'), as_of or MISSING),
        (tr('Stato', 'Status'), source.get('status') or 'data_missing'),
        (tr('Scadenza', 'Valid until'), source.get('valid_until') or MISSING),
        (tr('Nota provider', 'Provider note'), source.get('message') or MISSING),
    ), 5):
        _literal(ws, row, 2, name)
        _literal(ws, row, 4, value)


def _walk(value, path=(), inherited=None):
    """Yield raw leaves; period/metric labels are identifiers, never inferred dates."""
    inherited = inherited or {}
    if isinstance(value, dict):
        context = {**inherited, **{key: value[key] for key in _METADATA if key in value}}
        if value.get('table') is True and {'index', 'columns', 'data'} <= set(value):
            index, columns, cells = (value[key] for key in ('index', 'columns', 'data'))
            if (isinstance(index, list) and isinstance(columns, list) and isinstance(cells, list)
                    and len(index) == len(cells) and all(isinstance(line, list) and len(line) == len(columns) for line in cells)):
                for label, line in zip(index, cells):
                    for column, cell in zip(columns, line):
                        yield (*path, label, column), cell, context
                yield (*path, 'missing_cells'), value.get('missing_cells', MISSING), context
            else:
                yield (*path, 'malformed_table'), MISSING, context
        else:
            for key, child in value.items():
                if key not in _METADATA:
                    yield from _walk(child, (*path, key), context)
    elif isinstance(value, list):
        for index, child in enumerate(value):
            # List order remains visible. A period label is displayed in its own row.
            yield from _walk(child, (*path, index), inherited)
    else:
        yield path, value, inherited


def _display(value):
    if value is None:
        return MISSING
    if isinstance(value, float) and not isfinite(value):
        return MISSING
    return value


def _history(ws, source):
    data = source.get('data')
    row = 14
    if not isinstance(data, dict):
        _literal(ws, row, 2, tr('Bilanci assenti', 'Financial statements unavailable'))
        return row + 1
    errors = data.get('errors') if isinstance(data.get('errors'), dict) else {}
    for name in ('income_stmt', 'balance_sheet', 'cashflow'):
        _line(ws, row, name, 8, section=True)
        row += 1
        table = data.get(name)
        if not isinstance(table, dict) or table.get('table') is not True:
            _literal(ws, row, 2, MISSING)
            _literal(ws, row, 4, errors.get(name) or tr('Tabella non acquisita', 'Table not acquired'))
            row += 2
            continue
        index, columns, cells = (table.get(key) for key in ('index', 'columns', 'data'))
        if (not isinstance(index, list) or not isinstance(columns, list) or not isinstance(cells, list)
                or len(index) != len(cells) or any(not isinstance(line, list) or len(line) != len(columns) for line in cells)):
            _literal(ws, row, 2, tr('Tabella malformata', 'Malformed table'))
            row += 2
            continue
        _header(ws, row, [tr('Voce originale', 'Original item'), tr('Unita dichiarata', 'Declared unit'), *columns])
        row += 1
        unit = _unit(source, table)
        for label, line in zip(index, cells):
            _literal(ws, row, 2, label)
            _literal(ws, row, 3, unit)
            for column, value in enumerate(line, 4):
                _literal(ws, row, column, _display(value), AMOUNT if isinstance(value, (int, float)) and not isinstance(value, bool) else None)
            row += 1
        _literal(ws, row, 2, tr('Celle assenti dal provider', 'Provider missing cells'))
        _literal(ws, row, 4, table.get('missing_cells', MISSING))
        row += 2
    return row


def _structured(ws, source):
    _header(ws, 13, [tr('Campo / percorso provider', 'Provider field / path'),
                     tr('Valore grezzo', 'Raw value'), tr('Unita', 'Unit'),
                     tr('Fonte', 'Source'), tr('Data', 'Date'), tr('Stato', 'Status'),
                     tr('Scadenza', 'Valid until')])
    data = source.get('data')
    row = 14
    if data is None or data == {} or data == []:
        _literal(ws, row, 2, tr('Dati assenti', 'Data unavailable'))
        _literal(ws, row, 3, MISSING)
        return row + 1
    count = 0
    for path, value, context in _walk(data):
        _literal(ws, row, 2, '/'.join(map(str, path)) or 'value')
        _literal(ws, row, 3, _display(value), AMOUNT if isinstance(value, (int, float)) and not isinstance(value, bool) else None)
        _literal(ws, row, 4, context.get('unit') or _unit(source))
        _literal(ws, row, 5, context.get('source_id') or source.get('source_id') or MISSING)
        _literal(ws, row, 6, context.get('as_of') or source.get('as_of') or MISSING)
        _literal(ws, row, 7, context.get('status') or source.get('status') or 'data_missing')
        _literal(ws, row, 8, context.get('valid_until') or source.get('valid_until') or MISSING)
        row += 1
        count += 1
    if not count:
        _literal(ws, row, 2, tr('Nessun valore acquisito', 'No acquired values'))
        row += 1
    return row


def present_sourcebook(wb, payload):
    """Expose consumed records and preserve separately acquired raw provider views."""
    canonical = _present_consumed_sources(wb, payload)
    case = (payload.get('acquisition_snapshot') or {}).get('case')
    if not isinstance(case, dict):
        return canonical
    sources = case.get('sources') if isinstance(case.get('sources'), dict) else {}
    for name, title in (('financials', tr('Storico contabile', 'Financial history')),
                        ('consensus', 'Consensus'), ('guidance', 'Guidance')):
        source = sources.get(name) if isinstance(sources.get(name), dict) else {'status': 'data_missing'}
        ws = _sheet(wb, {'financials': 'Source History', 'consensus': 'Source Consensus',
                         'guidance': 'Source Guidance'}[name], title, 9)
        _line(ws, 4, tr('Dati grezzi dello snapshot: periodi e unita non normalizzati; nessuna comparabilita certificata.',
                        'Raw snapshot data: periods and units unnormalized; comparability is not certified.') +
                        (tr(' Fonti effettivamente consumate e motivazioni: Sources.',
                            ' Actually consumed sources and rationales: Sources.') if canonical else ''), 9, height=40)
        if canonical:
            ws['B4'].hyperlink = Hyperlink(ref='B4', location="'Sources'!B7")
        _metadata(ws, source, case.get('as_of'))
        last = _history(ws, source) if name == 'financials' else _structured(ws, source)
        ws.column_dimensions['B'].width = 58
        ws.column_dimensions['C'].width = 18 if name == 'financials' else 32
        if name == 'financials':
            for column in range(4, ws.max_column+1):
                ws.column_dimensions[get_column_letter(column)].width=24
        else:
            ws.column_dimensions['E'].width = 38
        for cells in ws.iter_rows(min_row=5):
            lines = 1
            for cell in cells:
                if cell.value is None:
                    continue
                numeric=isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool)
                cell.alignment = Alignment(vertical='top', wrap_text=not numeric,
                    horizontal='right' if numeric else 'left')
                width = ws.column_dimensions[get_column_letter(cell.column)].width
                lines = max(lines, sum(max(1, ceil(len(part) / max(1, width-2))) for part in str(cell.value).split('\n')))
            ws.row_dimensions[cells[0].row].height = min(409, max(22, 16*lines+6))
        ws.freeze_panes = 'D14'
        ws.print_title_rows = '13:13' if name != 'financials' else '2:4'
        _finish(ws, last, max(9, ws.max_column))
    return True
