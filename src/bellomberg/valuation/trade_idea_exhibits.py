"""Deterministic dossier exhibits bound to an exact common-model generation.

Numbers come from the calculation engine. The retained workbook baselines must
match them, and every displayed figure carries its actual workbook reference.
An empty Excel formula cache never turns into an invented chart value.
"""
from hashlib import sha256
from math import isclose, isfinite
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SCENARIOS = ('bear', 'base', 'bull')


def _finite(value):
    return type(value) in (int, float) and isfinite(value)


def _ref(sheet, cell):
    return "'" + sheet.replace("'", "''") + "'!" + cell


def _baselines(workbook):
    values = {}
    for sheet in workbook:
        if sheet.title not in ('Valuation', 'Engine Baseline', *SCENARIOS):
            continue
        for row in sheet:
            if len(row) < 2 or not isinstance(row[0].value, str) or not _finite(row[1].value):
                continue
            key = row[0].value if sheet.title in ('Valuation', 'Engine Baseline') else sheet.title + '.' + row[0].value
            values[key] = (row[1].value, _ref(sheet.title, row[1].coordinate))
    if 'Model Checks' in workbook:
        for row in workbook['Model Checks']:
            if len(row) >= 3 and row[1].value in SCENARIOS and _finite(row[2].value):
                values.setdefault(row[1].value + '.fair_value_per_share',
                                  (row[2].value, _ref('Model Checks', row[2].coordinate)))
    # Managed care preserves its income and capital tables transposed, while
    # the other intrinsic builders preserve a dotted-path raw scenario sheet.
    for label, prefix in (('Conto economico', 'rows'), ('Capitale e parent', 'capital.rows')):
        for scenario in SCENARIOS:
            name = label[:19] + ' ' + scenario
            if name not in workbook:
                continue
            for row in workbook[name]:
                if not row or not isinstance(row[0].value, str):
                    continue
                for index, cell in enumerate(row[1:]):
                    if _finite(cell.value):
                        values[f'{scenario}.{prefix}.{index}.{row[0].value}'] = (cell.value, _ref(name, cell.coordinate))
    return values


def model_exhibits(payload, *, workbook_path=None):
    """Return renderer-ready tables/charts or an explicit blocked report."""
    result = {key: payload.get(key) for key in ('snapshot_id', 'generation_id', 'valuation_date', 'workbook_sha256')}
    result.update(status='blocked', reasons=[], exhibits=[], contract='common_model_exhibits/1')
    result['headline_values'] = {}
    reasons, exhibits = result['reasons'], result['exhibits']
    path = Path(workbook_path or payload.get('path') or '')
    if not path.is_file() or not result['snapshot_id'] or not result['generation_id']:
        reasons.append('Exact workbook, snapshot and generation required')
        return result
    if sha256(path.read_bytes()).hexdigest() != result['workbook_sha256']:
        reasons.append('Workbook SHA-256 differs from the recorded generation')
        return result
    exposure = payload.get('method') == 'exposure_analysis'
    if not exposure and not payload.get('valuation_usability', {}).get('usable'):
        reasons.append('Economic model unusable; decisive value exhibits unavailable')
        return result
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        if exposure:
            from .trade_idea_exposure import exposure_exhibits
            return exposure_exhibits(payload, workbook, result)
        baselines = _baselines(workbook)
        cases = (payload.get('calculation_details', {}).get('scenarios') or
                 payload.get('managed_care', {}).get('scenarios') or {})
        def baseline(key, value):
            original = baselines.get(key)
            if not _finite(value) or not original or not isclose(value, original[0], rel_tol=1e-10, abs_tol=1e-9):
                raise ValueError('Engine/workbook baseline mismatch: ' + key)
            return original[1]

        def add(ident, title, kind, columns, rows, refs, *, unit, currency=None,
                period=None, scenario=None, explanation='', **extra):
            exhibits.append({'id': ident, 'title': title, 'kind': kind, 'columns': columns,
                'rows': rows, 'cell_refs': refs, 'unit': unit, 'currency': currency,
                'period': period or result['valuation_date'], 'scenario': scenario,
                'source': 'exact_common_engine_and_workbook', 'explanation': explanation, **extra})

        def bound_cell(sheet, coordinate):
            if sheet not in workbook or workbook[sheet][coordinate].value is None:
                raise ValueError('Decisive common workbook cell absent: ' + _ref(sheet, coordinate))
            return _ref(sheet, coordinate)

        # All families retain an engine baseline, even when the visible common
        # layout has different row/column conventions (e.g. a NAV snapshot).
        rows, refs = [], []
        nav = payload.get('method') in ('fund_nav', 'digital_asset_nav')
        for index, scenario in enumerate(SCENARIOS):
            raw = cases.get(scenario, {}).get('fair_value_per_share')
            raw_ref = baseline(scenario + '.fair_value_per_share', raw)
            quoted = payload.get('fair_value_' + scenario)
            if not _finite(quoted):
                raise ValueError('Quoted scenario value absent: ' + scenario)
            visible = bound_cell('Summary', f'C{8+index}' if nav else get_column_letter(4+index) + '8')
            rows.append([scenario, quoted, raw])
            refs.append([None, visible, raw_ref])
        add('scenarios', 'Scenari di valore', 'bar', ['Scenario', 'Valore quotato', 'Valore in valuta di bilancio'],
            rows, refs, unit='per share', currency=payload.get('currency'),
            column_units=[None, 'per share', 'per share'],
            column_currencies=[None, payload.get('currency'), payload.get('financial_currency')],
            explanation='Same engine, recorded valuation date and exact generation; no scenario probabilities.')
        result['headline_values']['fair_value_base'] = {'value': payload['fair_value_base'],
            'cell': refs[1][1], 'unit': 'per share', 'currency': payload.get('currency'),
            'period': result['valuation_date']}
        price = payload.get('price')
        quote = next((row for row in payload.get('acquisition_snapshot', {}).get('case', {}).get('records', [])
                      if row.get('scenario') == 'model' and row.get('driver') == 'quotation'), None)
        if quote and _finite(price) and quote['value'].get('price') == price:
            price_cell = None
            if payload.get('method') in ('operating_fcff', 'bank_residual_income'):
                price_cell = bound_cell('Summary', 'E18')
            elif not nav and 'Summary' in workbook:
                price_cell = bound_cell('Summary', 'E10')
            elif nav:
                bindings = payload.get('calculation_details', {}).get('workbook_bindings', {}).get('input_refs', [])
                price_cell = next((row['cell'] for row in bindings if row['scenario'] == 'model'
                                   and row['driver'] == 'quotation' and row['path'] == ['price']), None)
                # NAV's builder binds target inputs but not the quote in its
                # exported model link. The immutable raw quote remains visible
                # in the quality sheet; omit the headline if no binding exists.
            if price_cell:
                result['headline_values']['price'] = {'value': price, 'cell': price_cell,
                    'unit': 'per share', 'currency': payload.get('currency'),
                    'period': quote['value']['price_as_of']}

        base = cases.get('base', {})
        bridge = base.get('valuation_bridge') or {}
        if bridge:
            keys = ('pv_explicit_cash_flows', 'pv_terminal_value', 'enterprise_value', 'net_debt',
                    'equity_adjustments', 'equity_value')
            cells = {'pv_explicit_cash_flows': 'D48', 'pv_terminal_value': 'D46', 'enterprise_value': 'D49',
                     'net_debt': 'D50', 'equity_adjustments': 'D51', 'equity_value': 'D52'}
            rows, refs = [], []
            for key in keys:
                value = bridge.get(key)
                original = baseline('base.valuation_bridge.' + key, value)
                rows.append([key, value])
                refs.append([None, bound_cell('base', cells[key]) if payload.get('method') == 'operating_fcff' else original])
            add('valuation_bridge', 'Ponte al valore azionario', 'waterfall', ['Componente', 'Valore'], rows, refs,
                unit='million', currency=payload.get('financial_currency'), scenario='base',
                explanation='Explicit cash flows, terminal value and observed claims reconcile to equity value.')

        calendar = next((row.get('values') for row in payload.get('analytical_quality', {}).get('rows', [])
                         if row.get('scenario') == 'model' and row.get('driver') == 'calendar'), {}) or {}
        years = [period.get('end') for period in calendar.get('periods', [])]
        row_data = base.get('rows') or {}
        selected = [key for key in ('revenue', 'ebitda', 'ebit', 'ufcf') if key in row_data]
        if selected and years:
            rows, refs = [], []
            mapping = {'revenue': 8, 'ebitda': 14, 'ebit': 17, 'ufcf': 26}
            for index, year in enumerate(years):
                line, line_refs = [year], [None]
                for key in selected:
                    value = row_data[key][index]
                    original = baseline(f'base.rows.{key}.{index}', value)
                    line.append(value)
                    line_refs.append(bound_cell('base', get_column_letter(index+4) + str(mapping[key]))
                                     if payload.get('method') == 'operating_fcff' else original)
                rows.append(line); refs.append(line_refs)
            add('drivers', 'Driver operativi e flussi', 'line', ['Periodo', *selected], rows, refs,
                unit='million', currency=payload.get('financial_currency'), period=years,
                scenario='base', explanation='Calculated paths from the common engine, with linked worksheet cells.')

        if payload.get('method') == 'operating_fcff' and bridge:
            _operating_exhibits(payload, workbook, base, baseline, bound_cell, add)
        else:
            from .trade_idea_family_exhibits import family_exhibits
            family_exhibits(payload, base, baseline, add)
        if not any(row['id'] == 'valuation_bridge' for row in exhibits):
            # Method-specific decompositions are copied from the exact engine,
            # preserving their own labels and units rather than forcing an EV DCF.
            for name in ('bridge', 'components', 'nav_components', 'value_components'):
                components = base.get(name)
                if not isinstance(components, dict):
                    continue
                rows, refs = [], []
                for key, value in components.items():
                    if not _finite(value):
                        continue
                    rows.append([key, value]); refs.append([None, baseline('base.' + name + '.' + key, value)])
                if rows:
                    add('valuation_bridge', 'Ponte del metodo', 'waterfall', ['Componente', 'Valore'], rows, refs,
                        unit='method-native', currency=payload.get('financial_currency'), scenario='base')
                    break
        result['status'] = 'complete'
        present = {row['id'] for row in exhibits}
        result['exhibit_availability'] = {name: {'status': 'available' if name in present else 'unavailable',
            'reason': None if name in present else 'No exact common-engine/workbook exhibit for this method'}
            for name in ('scenarios', 'valuation_bridge', 'drivers', 'sensitivity', 'implicit_expectations')}
        if 'implicit_expectations' not in present and payload.get('method') == 'operating_fcff':
            market = payload.get('market_quote') or {}
            result['exhibit_availability']['implicit_expectations']['reason'] = (
                'Inverse model unavailable: current dated quotation is not usable; '
                + '; '.join(map(str, market.get('reasons') or [market.get('status') or 'quotation absent'])))
        result['unavailable_exhibits'] = [{'id': key, 'reason': value['reason']}
            for key, value in result['exhibit_availability'].items() if value['status'] != 'available']
        result['input_bindings'] = payload.get('calculation_details', {}).get('workbook_bindings', {}).get('input_refs', [])
    except (KeyError, IndexError, TypeError, ValueError) as exc:
        reasons.append(str(exc))
        result['exhibits'] = []
        result['headline_values'] = {}
    finally:
        workbook.close()
    return result


def _operating_exhibits(payload, workbook, base, baseline, bound_cell, add):
    """Use the shared DCF arithmetic also used by the workbook's sensitivity."""
    from .dcf_buyside_v3 import _dcf_value
    inputs = {(row['scenario'], row['driver']): row['values'] for row in payload['analytical_quality']['rows']}
    bridge = base['valuation_bridge']; terminal = base['terminal_bridge']
    q = inputs[('model', 'quotation')]
    factor = q['financial_to_quote_rate'] * q['quote_units_per_currency'] * q['shares_per_quote']
    spec = {'documented_inputs': True, 'discount_periods': bridge['discount_periods'], 'mid_year': False,
            'net_debt': bridge['net_debt'], 'shares_m': bridge['diluted_shares'],
            'diluted_shares_m': bridge['diluted_shares'],
            'equity_adjustments': [{'label': 'documented bridge', 'value_m': bridge['equity_adjustments']}]}
    rate, growth, ronic = base['wacc'], base['terminal_growth'], base['terminal_ronic']
    tax = inputs[('base', 'tax_rate')][-1]
    rows, refs = [], []
    for row_index, rate_shock in enumerate((-.01, -.005, 0., .005, .01), 8):
        shocked_rate = rate + rate_shock
        line, line_refs = [shocked_rate], [bound_cell('Sensitivity', 'B' + str(row_index))]
        for column, growth_shock in enumerate((-.005, -.0025, 0., .0025, .005), 3):
            shocked_growth = growth + growth_shock
            value = None
            if shocked_rate > max(shocked_growth, 0.) and ronic > max(shocked_growth, 0.) and shocked_growth > -1:
                value = _dcf_value(spec, base['rows']['ufcf'], shocked_rate, shocked_growth,
                    ebit_terminal=terminal['normalized_ebit'], ronic=ronic, tax_term=tax) * factor
            line.append(value)
            line_refs.append(bound_cell('Sensitivity', get_column_letter(column) + str(row_index)))
        rows.append(line); refs.append(line_refs)
    add('sensitivity', 'Sensibilita WACC / crescita stabile', 'sensitivity',
        ['WACC', *[growth+x for x in (-.005, -.0025, 0., .0025, .005)]], rows, refs,
        unit='per share', currency=payload.get('currency'), scenario='base',
        explanation='Same explicit cash flows, terminal EBIT, taxes, reinvestment and claims; invalid combinations are missing.')
    market = payload.get('market_quote') or {}
    price = market.get('price')
    if market.get('status') == 'ok' and _finite(price) and price > 0:
        denominator = (1+growth) * (1-tax) * (1-growth/ronic)
        implied = ((price/factor*bridge['diluted_shares'] + bridge['net_debt'] - bridge['equity_adjustments']
            - bridge['pv_explicit_cash_flows']) * (1+rate)**bridge['discount_periods'][-1] * (rate-growth)
            / denominator) if denominator > 0 and rate > growth else None
        if _finite(implied):
            add('implicit_expectations', 'EBIT stabile implicito nel prezzo osservato', 'bar',
                ['Misura', 'EBIT'], [['EBIT normalizzato del modello', terminal['normalized_ebit']],
                                   ['EBIT stabile implicito', implied]],
                [[None, bound_cell('Sensitivity', 'D19')], [None, bound_cell('Sensitivity', 'D20')]],
                unit='million', currency=payload.get('financial_currency'), scenario='base',
                explanation='Conditional on the base model and dated usable quote; an inverse model, never consensus.')
