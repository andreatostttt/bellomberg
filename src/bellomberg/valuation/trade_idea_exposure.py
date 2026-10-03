"""Observed exposure analysis through the common acquisition and workbook path.

There is no corporate fair value, target price or substitute DCF. Completeness of
an exposure analysis has a separate contract from valuation usability.
"""
from copy import deepcopy
from hashlib import sha256
from math import isclose, isfinite
from pathlib import Path
import re


SCHEMA = {
    'perimeter': ('instrument_identity', 'contract', 'scope', 'opening', 'contract', 'model'),
    'calendar': ('exposure_terms', 'contract', 'calendar', 'opening', 'contract', 'model'),
    'identity': ('instrument_identity', 'contract', 'instrument', 'opening', 'contract', 'model'),
    'holdings': ('exposure_terms', 'contract', 'net_asset_exposure', 'opening', 'contract', 'model'),
    'terms': ('exposure_terms', 'contract', 'replication', 'opening', 'contract', 'model'),
    'fees': ('fees_and_risks', 'contract', 'annual_cost', 'opening', 'contract', 'model'),
    'risks': ('fees_and_risks', 'contract', 'risk_disclosure', 'opening', 'contract', 'model'),
    'liquidity': ('fees_and_risks', 'contract', 'market_liquidity', 'opening', 'contract', 'model'),
    'quotation': ('quotation_units', 'contract', 'quotation', 'opening', 'contract', 'model'),
}
PREPARATION_POLICY = (
    'Exposure snapshot, no corporate valuation. model only, scenarios={}, analysis_rationale required. '
    'Every identity/holdings/terms/fees/risks/liquidity/quotation value is an observed source contract '
    'with historical record_pointer proof. Perimeter/calendar are cited structural modeling choices. '
    'holdings={positions:[{id,name,weight,asset_class,country,currency}],basis:net_asset_exposure}; '
    'terms={replication,benchmark,gross_exposure,net_exposure,leverage,effective_date}; '
    'fees={annual_expense_ratio,other_annual_fees_ratio}; risks={issuer,counterparty,custody,liquidity,tracking}; '
    'liquidity={daily_volume,volume_unit,spread_ratio,observed_date,provider}; '
    'identity={ticker,name,instrument,issuer,legal_structure,underlying}. '
    'Position weights are signed fractions of NAV, not percentages or normalized index composition. '
    'All positions must reconcile to exact declared gross/net exposure and leverage. '
    'Explicit zeros require source observations, missing costs/holdings/risks remain missing.'
)


def _finite(value):
    return type(value) in (int, float) and isfinite(value)


def project_exposure(values, *, ticker, instrument, valuation_date, cutoff=None):
    """Deterministic reconciliations; no statistical/return assumptions are added."""
    from .trade_idea_commodity_exposure import selected, project
    if selected(values):
        return project(values, ticker=ticker, instrument=instrument, cutoff=cutoff or valuation_date)
    issues = []
    def exact(name, keys):
        value = values.get(name)
        if not isinstance(value, dict) or set(value) != set(keys):
            issues.append(name + ': exact observed contract required')
            return {}
        return value
    identity = exact('identity', ('ticker', 'name', 'instrument', 'issuer', 'legal_structure', 'underlying'))
    if identity and (identity['ticker'] != ticker or identity['instrument'] != instrument or
                     any(not isinstance(value, str) or not value.strip() for value in identity.values())):
        issues.append('identity: instrument, ticker and legal identity differ from the acquired instrument')
    holdings = exact('holdings', ('positions', 'basis'))
    positions = holdings.get('positions')
    if holdings.get('basis') != 'net_asset_exposure' or not isinstance(positions, list) or not positions:
        issues.append('holdings: complete NAV exposure positions required')
        positions = []
    ids = set()
    for position in positions:
        if (not isinstance(position, dict) or set(position) != {'id', 'name', 'weight', 'asset_class', 'country', 'currency'}
                or any(not isinstance(position.get(key), str) or not position[key].strip()
                       for key in ('id', 'name', 'asset_class', 'country', 'currency'))
                or not _finite(position.get('weight')) or position.get('id') in ids):
            issues.append('holdings: missing/duplicate position identity, classifications or signed weight')
        else:
            ids.add(position['id'])
    terms = exact('terms', ('replication', 'benchmark', 'gross_exposure', 'net_exposure', 'leverage', 'effective_date'))
    if terms:
        if (terms['effective_date'] != valuation_date or terms['replication'] not in
                ('physical_full', 'physical_sampled', 'synthetic', 'direct_asset')
                or not isinstance(terms['benchmark'], str) or not terms['benchmark'].strip()
                or any(not _finite(terms[key]) for key in ('gross_exposure', 'net_exposure', 'leverage'))
                or terms['gross_exposure'] < abs(terms['net_exposure']) or terms['leverage'] < 0):
            issues.append('terms: observed replication, benchmark, dated gross/net exposure and leverage required')
    fees = exact('fees', ('annual_expense_ratio', 'other_annual_fees_ratio'))
    if fees and any(not _finite(value) or not 0 <= value <= 1 for value in fees.values()):
        issues.append('fees: both explicit annual cost fractions must be observed and in domain')
    risks = exact('risks', ('issuer', 'counterparty', 'custody', 'liquidity', 'tracking'))
    if risks and any(not isinstance(value, str) or not value.strip() for value in risks.values()):
        issues.append('risks: exact primary risk disclosures required')
    liquidity = exact('liquidity', ('daily_volume', 'volume_unit', 'spread_ratio', 'observed_date', 'provider'))
    if liquidity and (not _finite(liquidity['daily_volume']) or liquidity['daily_volume'] < 0
            or liquidity['volume_unit'] != 'quoted units' or not _finite(liquidity['spread_ratio'])
            or not 0 <= liquidity['spread_ratio'] <= 1 or liquidity['observed_date'] != valuation_date
            or not isinstance(liquidity['provider'], str) or not liquidity['provider'].strip()):
        issues.append('liquidity: dated volume, unit, spread and identified provider required')
    if issues:
        return {'status': 'incomplete', 'issues': issues}
    gross = sum(abs(position['weight']) for position in positions)
    net = sum(position['weight'] for position in positions)
    if not isclose(gross, terms['gross_exposure'], abs_tol=1e-9) or not isclose(net, terms['net_exposure'], abs_tol=1e-9):
        issues.append('holdings: weights do not reconcile to declared gross/net exposure; omitted positions are not zero')
    if not isclose(gross, terms['leverage'], abs_tol=1e-9):
        issues.append('leverage: declared NAV exposure leverage does not reconcile to the holdings')
    buckets = {key: {} for key in ('asset_class', 'country', 'currency')}
    for position in positions:
        for key, bucket in buckets.items():
            bucket[position[key]] = bucket.get(position[key], 0.) + position['weight']
    return {'status': 'incomplete' if issues else 'complete', 'issues': issues,
        'positions': deepcopy(positions), 'identity': deepcopy(identity), 'terms': deepcopy(terms),
        'fees': deepcopy(fees), 'risks': deepcopy(risks), 'liquidity': deepcopy(liquidity),
        'gross_exposure': gross, 'net_exposure': net,
        'largest_absolute_weight': max(abs(position['weight']) for position in positions),
        'squared_weight_sum': sum(position['weight']**2 for position in positions),
        'annual_declared_cost_ratio': sum(fees.values()),
        'buckets': {kind: [{'label': label, 'weight': weight} for label, weight in bucket.items()]
                    for kind, bucket in buckets.items()},
        'concentration_basis': 'Signed NAV exposure weights; squared-weight sum is not an effective holding count for leveraged or short exposure.'}


def generate_exposure(bundle, *, output_dir, metadata):
    from .documented_inputs import bind_inputs
    from .dcf_engine import _write_payload_sidecar, REPORT_DIR
    from .trade_idea_commodity_exposure import selected_records, bind as commodity_bind, CONTRACT
    commodity = selected_records(bundle['case']['records'])
    bound = commodity_bind(bundle) if commodity else bind_inputs(bundle, SCHEMA, horizon='snapshot', scenario_names=())
    values = bound['values']['model']
    projection = project_exposure(values, ticker=bundle['case']['ticker'],
        instrument=bundle['decision']['instrument'], valuation_date=bound['calendar'].get('valuation_date'), cutoff=bundle['case']['as_of']) if not bound['issues'] else {'status': 'incomplete', 'issues': []}
    reasons = [row['reason'] for row in bound['issues']] + projection['issues']
    complete = not reasons and projection['status'] == 'complete'
    from .market_quote import build_market_quote, _percent
    from .trade_idea_model import _current_quotation_evidence
    quotation = bound['quotation']
    market_quote = build_market_quote(bundle, quotation, {})
    quote_problem, quote_receipt = _current_quotation_evidence(
        bundle, {'model': {'quotation': {'value': quotation}}}, bundle['case']['as_of'])
    market_quote.update(quote_receipt)
    market_quote.update(status='source_unavailable' if quote_problem else 'ok',
        message=quote_problem or 'Quotazione regularMarketPrice osservata con ricevuta verificata; '
            'snapshot contabile e prezzo pubblicato dall\'emittente restano alle proprie date. Fair value non applicabile.',
        fv_basis='not_applicable_observed_exposure', intrinsic_value_applicable=False,
        comparison_status='not_applicable', comparison_qualified=False,
        comparison_reason='Observed exposure analysis has no intrinsic fair-value comparison',
        price_move_since_valuation_pct=None if quote_problem else _percent(market_quote['price'], quotation.get('price')))
    result = {**metadata, 'ticker': bundle['case']['ticker'], 'engine': 'exposure_analysis', 'method': 'exposure_analysis',
        'ok': complete, 'valuation_date': bound['calendar'].get('valuation_date'),
        'currency': bound['quotation'].get('quote_unit'), 'financial_currency': bound['perimeter'].get('currency'),
        'price': bound['quotation'].get('price'), 'market_quote': market_quote,
        'valuation_basis': 'Observed exposure snapshot; corporate fair value is not applicable.',
        'exposure_analysis': projection, 'intrinsic_value_applicable': False,
        'analysis_usability': {'usable': complete, 'reasons': reasons, 'contract': CONTRACT if commodity else 'documented_exposure/1',
                              'intrinsic_value_applicable': False},
        'analytical_quality': {'method_id': 'exposure_analysis', 'status': 'DOCUMENTATA' if complete else 'INCOMPLETA',
                              'issues': reasons, 'rows': bound['evidence'], 'snapshot': {'forecast_years': []}},
        'input_consumption': {'status': 'complete' if complete else 'incomplete',
            'consumed_fields': sorted({row['field'] for row in bound['consumed']}),
            'unconsumed_fields': [] if complete else sorted({row['field'] for row in bound['issues']}),
            'consumed_records': bound['consumed']},
        'sanity': {'status': 'not_applicable', 'severity': 'OK' if complete else 'BLOCK',
                   'method_id': 'exposure_analysis', 'intrinsic_value_applicable': False},
        'calculation_details': {'exposure': projection}, 'acquisition_tasks': deepcopy(bundle['acquisition_tasks']) + bound['issues']}
    if complete:
        result['path'] = build_exposure_workbook(result, output_dir or REPORT_DIR)
        return _write_payload_sidecar(result)
    result['error'] = 'Exposure analysis incomplete: ' + '; '.join(reasons)
    return result


def build_exposure_workbook(payload, output_dir):
    from .trade_idea_commodity_exposure import CONTRACT, workbook
    if payload.get('analysis_usability', {}).get('contract') == CONTRACT:
        return workbook(payload, output_dir)
    from openpyxl import Workbook
    from openpyxl.workbook.properties import CalcProperties
    from .dcf_quality_sheet import append_quality_sheet, append_sector_quality_sheet
    from .sourcebook_presentation import present_sourcebook
    from .documented_presentation import _sheet, _line, _header, _put, _finish, PERCENT, PRICE
    from .documented_formulas import _formula
    directory = Path(output_dir); directory.mkdir(parents=True, exist_ok=True)
    path = directory / ('EXPOSURE_' + re.sub(r'[^A-Za-z0-9_-]', '_', payload['ticker']) + '_' + payload['generation_id'] + '.xlsx')
    wb = Workbook(); wb.remove(wb.active)
    p = payload['exposure_analysis']
    ws = _sheet(wb, 'Exposure', payload['ticker'] + ' | Analisi delle esposizioni', 8)
    _line(ws, 4, 'Snapshot osservazionale: il fair value aziendale non è applicabile. Pesi riferiti al NAV e dati osservati alla data dichiarata.', 8, height=42)
    _header(ws, 7, ['ID', 'Nome', 'Peso NAV', 'Classe', 'Paese', 'Valuta', 'Peso assoluto'])
    refs = []
    for index, row in enumerate(p['positions'], 8):
        for column, key in enumerate(('id', 'name', 'weight', 'asset_class', 'country', 'currency'), 2):
            _put(ws, index, column, row[key], PERCENT if key == 'weight' else None)
        _formula(ws, index, 8, f'ABS(D{index})', PERCENT)
        refs.append({'scenario': 'model', 'driver': 'holdings', 'path': ['positions', index-8, 'weight'], 'cell': "'Exposure'!D" + str(index)})
    end = 7 + len(p['positions']); _finish(ws, end+2, 8)
    summary = _sheet(wb, 'Summary', payload['ticker'] + ' | Esposizione osservata', 6)
    _line(summary, 4, 'Prezzo, costi e portafoglio sono osservazioni. Nessun target price, DCF o upside intrinseco.', 6, height=38)
    _header(summary, 7, ['Misura', 'Unità', 'Valore'])
    expressions = [('Esposizione lorda', 'ratio', f'SUM(\'Exposure\'!H8:H{end})', p['gross_exposure']),
        ('Esposizione netta', 'ratio', f'SUM(\'Exposure\'!D8:D{end})', p['net_exposure']),
        ('Massimo peso assoluto', 'ratio', f'MAX(\'Exposure\'!H8:H{end})', p['largest_absolute_weight']),
        ('Somma pesi al quadrato', 'ratio', f'SUMSQ(\'Exposure\'!D8:D{end})', p['squared_weight_sum'])]
    baseline = wb.create_sheet('Engine Baseline'); baseline.append(['Engine path', 'Recorded value'])
    metrics = {}
    for row, (label, unit, expression, value) in enumerate(expressions, 8):
        _put(summary, row, 2, label); _put(summary, row, 3, unit); _formula(summary, row, 4, expression, PERCENT)
        key = ('gross_exposure', 'net_exposure', 'largest_absolute_weight', 'squared_weight_sum')[row-8]
        baseline.append([key, value]); metrics[key] = "'Summary'!D" + str(row)
    _put(summary, 13, 2, 'Costi annuali dichiarati'); _put(summary, 13, 3, 'ratio')
    baseline.append(['annual_declared_cost_ratio', p['annual_declared_cost_ratio']])
    metrics['annual_declared_cost_ratio'] = "'Summary'!D13"
    _put(summary, 15, 2, 'Prezzo osservato'); _put(summary, 15, 3, payload['currency']); _put(summary, 15, 4, payload['price'], PRICE)
    _put(summary, 17, 2, 'Data osservazioni'); _put(summary, 17, 4, payload['valuation_date'])
    _put(summary, 18, 2, 'Replica'); _put(summary, 18, 4, p['terms']['replication'])
    _put(summary, 19, 2, 'Benchmark'); _put(summary, 19, 4, p['terms']['benchmark']); _finish(summary, 21, 6)
    risks = _sheet(wb, 'Terms and risks', 'Termini, costi e rischi pubblicati', 6)
    _header(risks, 7, ['Dato', 'Valore'])
    fee_cells = []
    row = 8
    for row, name, key, value in _observation_rows(p):
        _put(risks, row, 2, name + '.' + key); _put(risks, row, 3, value)
        if _finite(value):
            cell = "'Terms and risks'!C" + str(row)
            refs.append({'scenario': 'model', 'driver': name, 'path': [key], 'cell': cell})
            if name == 'fees':
                fee_cells.append(cell)
    _formula(summary, 13, 4, '+'.join(fee_cells), PERCENT)
    _finish(risks, row+2, 6)
    append_quality_sheet(wb, payload['analytical_quality']); append_sector_quality_sheet(wb, payload)
    present_sourcebook(wb, payload)
    wb.move_sheet(summary, -wb.index(summary)); wb.active = 0
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    payload['calculation_details']['workbook_bindings'] = {'contract': 'common_workbook_bindings/1',
        'input_refs': refs, 'metrics': metrics, 'price': "'Summary'!D15"}
    wb.save(path); wb.close()
    return str(path)


def _observation_rows(projection):
    # Source dictionaries may be serialized with sorted keys in durable stores.
    # Workbook rows are an explicit layout contract, never insertion order.
    fields = {'identity': ('ticker', 'name', 'instrument', 'issuer', 'legal_structure', 'underlying'),
        'terms': ('replication', 'benchmark', 'gross_exposure', 'net_exposure', 'leverage', 'effective_date'),
        'fees': ('annual_expense_ratio', 'other_annual_fees_ratio'),
        'liquidity': ('daily_volume', 'volume_unit', 'spread_ratio', 'observed_date', 'provider'),
        'risks': ('issuer', 'counterparty', 'custody', 'liquidity', 'tracking')}
    row = 8
    for name in ('identity', 'terms', 'fees', 'liquidity', 'risks'):
        for key in fields[name]:
            yield row, name, key, projection[name][key]
            row += 1


def _verify_exposure_workbook(payload, workbook, projection, packet):
    """Compare every observation and the live reconciliation formula contract."""
    if packet.get('contract') != 'common_workbook_bindings/1':
        raise ValueError('Exposure workbook binding contract missing')
    expected_refs = []
    ws = workbook['Exposure']
    for row_index, position in enumerate(projection['positions'], 8):
        for column, key in enumerate(('id', 'name', 'weight', 'asset_class', 'country', 'currency'), 2):
            if ws.cell(row_index, column).value != position[key]:
                raise ValueError('Exposure workbook observation mismatch: holdings/' + str(row_index-8) + '/' + key)
        if ws.cell(row_index, 8).value != '=ABS(D' + str(row_index) + ')':
            raise ValueError('Exposure absolute-weight formula differs from acquired holdings')
        expected_refs.append({'scenario': 'model', 'driver': 'holdings',
            'path': ['positions', row_index-8, 'weight'], 'cell': "'Exposure'!D" + str(row_index)})
    end = 7 + len(projection['positions'])
    # Extra observations must not change the visible/formula perimeter.
    if any(ws.cell(end+1, column).value is not None for column in range(2, 9)):
        raise ValueError('Exposure workbook contains an unbound extra holding')
    fees = []
    terms = workbook['Terms and risks']
    for row, name, key, value in _observation_rows(projection):
        if terms.cell(row, 2).value != name + '.' + key or terms.cell(row, 3).value != value:
            raise ValueError('Exposure workbook observation mismatch: ' + name + '/' + key)
        if _finite(value):
            cell = "'Terms and risks'!C" + str(row)
            expected_refs.append({'scenario': 'model', 'driver': name, 'path': [key], 'cell': cell})
            if name == 'fees':
                fees.append(cell)
    if packet.get('input_refs') != expected_refs:
        raise ValueError('Exposure workbook input bindings differ from exact observations')
    formulas = {'gross_exposure': (8, f"SUM('Exposure'!H8:H{end})"),
        'net_exposure': (9, f"SUM('Exposure'!D8:D{end})"),
        'largest_absolute_weight': (10, f"MAX('Exposure'!H8:H{end})"),
        'squared_weight_sum': (11, f"SUMSQ('Exposure'!D8:D{end})"),
        'annual_declared_cost_ratio': (13, '+'.join(fees))}
    expected_metrics = {key: "'Summary'!D" + str(row) for key, (row, _) in formulas.items()}
    if packet.get('metrics') != expected_metrics or packet.get('price') != "'Summary'!D15":
        raise ValueError('Exposure decisive workbook bindings differ from the common layout')
    summary = workbook['Summary']
    for key, (row, formula) in formulas.items():
        if summary.cell(row, 4).value != '=' + formula:
            raise ValueError('Exposure reconciliation formula mismatch: ' + key)
    for row, value in ((15, payload['price']), (17, payload['valuation_date']),
                       (18, projection['terms']['replication']), (19, projection['terms']['benchmark'])):
        if summary.cell(row, 4).value != value:
            raise ValueError('Exposure summary observation mismatch at D' + str(row))


def exposure_exhibits(payload, workbook, result):
    """Verify and bind the observational packet, without valuation usability."""
    p = payload.get('exposure_analysis') or {}
    if payload.get('analysis_usability', {}).get('usable') is not True or p.get('status') != 'complete':
        raise ValueError('Exact exposure analysis is incomplete')
    from .sector_analysis import validate_bundle
    bundle = validate_bundle(payload.get('acquisition_snapshot'), payload.get('ticker'))
    if bundle['snapshot_id'] != payload.get('snapshot_id') or bundle['decision'] != payload.get('valuation_decision'):
        raise ValueError('Exposure acquisition/generation identity differs')
    recomputed = project_exposure({row['driver']: row['value'] for row in bundle['case']['records']},
        ticker=bundle['case']['ticker'], instrument=bundle['decision']['instrument'], valuation_date=payload['valuation_date'], cutoff=bundle['case']['as_of'])
    if recomputed != p:
        raise ValueError('Exposure calculation differs from its exact acquired observations')
    from .trade_idea_commodity_exposure import CONTRACT, exhibits as commodity_exhibits
    if payload.get('analysis_usability', {}).get('contract') == CONTRACT:
        return commodity_exhibits(payload, workbook, result)
    baseline = {row[0].value: row[1].value for row in workbook['Engine Baseline'] if len(row) >= 2}
    packet = payload['calculation_details']['workbook_bindings']
    _verify_exposure_workbook(payload, workbook, p, packet)
    for key, cell in packet['metrics'].items():
        if baseline.get(key) != p[key]:
            raise ValueError('Exposure engine/workbook baseline mismatch: ' + key)
    metric_order = ('gross_exposure', 'net_exposure', 'largest_absolute_weight',
                    'squared_weight_sum', 'annual_declared_cost_ratio')
    exhibits = [{'id': 'exposure_metrics', 'title': 'Esposizione e concentrazione', 'kind': 'table',
        'columns': ['Misura', 'Valore'], 'rows': [[key, p[key]] for key in metric_order],
        'cell_refs': [[None, packet['metrics'][key]] for key in metric_order], 'unit': 'ratio',
        'currency': None, 'period': payload['valuation_date'], 'scenario': 'observed',
        'source': 'exact_common_engine_and_workbook', 'explanation': p['concentration_basis']},
        {'id': 'holdings', 'title': 'Esposizioni osservate', 'kind': 'bar', 'columns': ['Strumento', 'Peso NAV'],
         'rows': [[row['name'], row['weight']] for row in p['positions']],
         'cell_refs': [[None, "'Exposure'!D" + str(index)] for index in range(8, 8+len(p['positions']))],
         'unit': 'ratio', 'currency': None, 'period': payload['valuation_date'], 'scenario': 'observed',
         'source': 'exact_common_engine_and_workbook', 'explanation': 'Signed fractions of NAV; no valuation scenarios or probabilities.'}]
    result.update(status='complete', exhibits=exhibits, input_bindings=deepcopy(packet['input_refs']),
        unavailable_exhibits=[{'id': key, 'reason': 'Corporate fair value is not applicable to this observed exposure analysis'}
                              for key in ('valuation_bridge', 'scenarios', 'sensitivity', 'implicit_expectations')],
        intrinsic_value_applicable=False, analysis_usability=deepcopy(payload['analysis_usability']))
    result['headline_values'] = {'price': {'value': payload['price'], 'cell': packet['price'], 'unit': 'per quoted unit',
                                         'currency': payload['currency'], 'period': payload['valuation_date']}}
    return result
